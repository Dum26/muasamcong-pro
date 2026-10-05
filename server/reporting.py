"""Whitelisted, user-scoped reporting over observed public procurement data."""
from __future__ import annotations

import csv
import io
import os
import re
import time
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, case, func, insert, or_, select, update, delete

from server.database import (contractors, get_db, investors, now, plans,
    report_runs, saved_reports, tender_changes, tenders, user_dashboards)

MAX_PAGE_SIZE = 100
MAX_EXPORT_ROWS = 5000


def export_sync_max_rows() -> int:
    try:
        return max(100, min(int(os.environ.get("EXPORT_SYNC_MAX_ROWS", MAX_EXPORT_ROWS)), MAX_EXPORT_ROWS))
    except ValueError:
        return MAX_EXPORT_ROWS
FILTER_OPS = {"equals", "not_equals", "contains", "starts_with", "greater_than",
              "less_than", "at_least", "at_most", "between", "is_empty",
              "is_not_empty", "before", "after"}
NUMERIC_OPS = {"count", "sum", "avg", "min", "max"}
TEXT_OPS = {"count", "min", "max"}


@dataclass(frozen=True)
class ReportDataset:
    name: str
    table: object
    fields: dict[str, object]
    numeric: frozenset[str]
    dates: frozenset[str]
    identity: str
    default_columns: tuple[str, ...]
    base_condition: object | None = None


TENDER_FIELDS = {key: tenders.c[key] for key in (
    "tender_code", "khlcnt_code", "title", "investor_id", "investor_name",
    "procuring_entity_name", "contractor_id", "winning_contractor",
    "package_price", "winning_price", "publish_date", "bid_closing_date",
    "result_date", "status", "selection_method", "contract_type", "location")}
PLAN_FIELDS = {key: plans.c[key] for key in (
    "plan_code", "title", "investor_id", "investor_name", "package_price", "publish_date")}
CONTRACTOR_FIELDS = {key: contractors.c[key] for key in (
    "id", "name", "tax_code", "first_seen_at", "last_seen_at")}
INVESTOR_FIELDS = {key: investors.c[key] for key in (
    "id", "name", "identifier", "first_seen_at", "last_seen_at")}
CHANGE_FIELDS = {key: tender_changes.c[key] for key in (
    "id", "tender_code", "field_name", "detected_at")}

DATASETS = {
    "tenders": ReportDataset("tenders", tenders, TENDER_FIELDS,
        frozenset({"package_price", "winning_price"}),
        frozenset({"publish_date", "bid_closing_date", "result_date"}),
        "tender_code", ("tender_code", "title", "investor_name", "package_price", "publish_date")),
    "khlcnt": ReportDataset("khlcnt", plans, PLAN_FIELDS,
        frozenset({"package_price"}), frozenset({"publish_date"}),
        "plan_code", ("plan_code", "title", "investor_name", "publish_date")),
    "kqlcnt": ReportDataset("kqlcnt", tenders, TENDER_FIELDS,
        frozenset({"package_price", "winning_price"}),
        frozenset({"publish_date", "bid_closing_date", "result_date"}),
        "tender_code", ("tender_code", "title", "winning_contractor", "winning_price", "result_date"),
        tenders.c.winning_contractor.is_not(None)),
    "contractors": ReportDataset("contractors", contractors, CONTRACTOR_FIELDS,
        frozenset(), frozenset({"first_seen_at", "last_seen_at"}),
        "id", ("name", "tax_code", "last_seen_at")),
    "investors": ReportDataset("investors", investors, INVESTOR_FIELDS,
        frozenset(), frozenset({"first_seen_at", "last_seen_at"}),
        "id", ("name", "identifier", "last_seen_at")),
    "changes": ReportDataset("changes", tender_changes, CHANGE_FIELDS,
        frozenset(), frozenset({"detected_at"}),
        "id", ("tender_code", "field_name", "detected_at")),
    "deadlines": ReportDataset("deadlines", tenders, TENDER_FIELDS,
        frozenset({"package_price", "winning_price"}),
        frozenset({"publish_date", "bid_closing_date", "result_date"}),
        "tender_code", ("tender_code", "title", "bid_closing_date", "status"),
        tenders.c.bid_closing_date.is_not(None)),
}


def _id() -> str:
    return uuid.uuid4().hex


def _safe_value(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date,)):
        return value.isoformat()
    return value


def _date(value) -> str:
    value = str(value or "")
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise ValueError("Ngày lọc phải có dạng YYYY-MM-DD.") from exc


def _number(value) -> Decimal:
    try:
        result = Decimal(str(value))
        if result.is_finite():
            return result
    except (InvalidOperation, TypeError):
        pass
    raise ValueError("Giá trị lọc số không hợp lệ.")


def _operand(dataset: ReportDataset, field: str, value):
    if field in dataset.dates:
        return _date(value)
    if field in dataset.numeric:
        return _number(value)
    result = str(value or "").strip()
    if not result or len(result) > 320:
        raise ValueError("Giá trị bộ lọc phải có từ 1 đến 320 ký tự.")
    return result


def _conditions(dataset: ReportDataset, filters: list) -> list:
    if not isinstance(filters, list) or len(filters) > 12:
        raise ValueError("Tối đa 12 điều kiện lọc.")
    conditions = [dataset.base_condition] if dataset.base_condition is not None else []
    for item in filters:
        if not isinstance(item, dict):
            raise ValueError("Điều kiện lọc không hợp lệ.")
        field, op = item.get("field"), item.get("operator")
        if field not in dataset.fields or op not in FILTER_OPS:
            raise ValueError("Trường hoặc phép lọc không được hỗ trợ.")
        column = dataset.fields[field]
        if op == "is_empty":
            conditions.append(or_(column.is_(None), column == "") if field not in dataset.numeric else column.is_(None))
            continue
        if op == "is_not_empty":
            conditions.append(and_(column.is_not(None), column != "") if field not in dataset.numeric else column.is_not(None))
            continue
        if op == "between":
            values = item.get("value")
            if not isinstance(values, list) or len(values) != 2:
                raise ValueError("Bộ lọc between cần hai giá trị.")
            low, high = (_operand(dataset, field, value) for value in values)
            if low > high:
                raise ValueError("Khoảng lọc bị đảo chiều.")
            if field in dataset.dates:
                upper = date.fromisoformat(high).toordinal() + 1
                conditions.append(and_(column >= low, column < date.fromordinal(upper).isoformat()))
            else:
                conditions.append(column.between(low, high))
            continue
        value = _operand(dataset, field, item.get("value"))
        if op in {"contains", "starts_with"}:
            if field in dataset.numeric or field in dataset.dates:
                raise ValueError("Chỉ có thể tìm chuỗi trong trường văn bản.")
            escaped = value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            pattern = escaped + "%" if op == "starts_with" else "%" + escaped + "%"
            conditions.append(column.ilike(pattern, escape="\\"))
        elif op == "equals":
            conditions.append(column == value)
        elif op == "not_equals":
            conditions.append(column != value)
        elif op in {"before", "after"}:
            if field not in dataset.dates:
                raise ValueError("before/after chỉ dùng cho ngày.")
            conditions.append(column < value if op == "before" else column >= value)
        elif op in {"greater_than", "less_than", "at_least", "at_most"}:
            if field not in dataset.numeric and field not in dataset.dates:
                raise ValueError("So sánh lớn/nhỏ chỉ dùng cho số hoặc ngày.")
            conditions.append({"greater_than": column > value, "less_than": column < value,
                "at_least": column >= value, "at_most": column <= value}[op])
    return conditions


def _group(dataset: ReportDataset, group_by: str):
    if not group_by:
        return None
    if group_by == "all":
        return func.substr(dataset.fields[dataset.identity], 1, 0).label("group")
    if group_by in dataset.fields and group_by not in dataset.numeric:
        return dataset.fields[group_by].label("group")
    for prefix, length in (("day", 10), ("month", 7), ("year", 4)):
        if group_by.startswith(prefix + ":"):
            field = group_by.split(":", 1)[1]
            if field not in dataset.dates:
                break
            return func.substr(dataset.fields[field], 1, length).label("group")
    if group_by.startswith("quarter:"):
        field = group_by.split(":", 1)[1]
        if field in dataset.dates:
            month = func.substr(dataset.fields[field], 6, 2)
            quarter = case((month <= "03", "Q1"), (month <= "06", "Q2"),
                           (month <= "09", "Q3"), else_="Q4")
            return (func.substr(dataset.fields[field], 1, 4) + "-" + quarter).label("group")
    raise ValueError("Nhóm dữ liệu không được hỗ trợ.")


def validate_spec(spec: dict) -> dict:
    if not isinstance(spec, dict) or spec.get("dataset") not in DATASETS:
        raise ValueError("Dataset không được hỗ trợ.")
    dataset = DATASETS[spec["dataset"]]
    columns = spec.get("columns") or list(dataset.default_columns)
    if not isinstance(columns, list) or not 1 <= len(columns) <= 16 or any(
            not isinstance(field, str) or field not in dataset.fields for field in columns):
        raise ValueError("Cột báo cáo không hợp lệ.")
    columns = list(dict.fromkeys(columns))
    filters = spec.get("filters") or []
    _conditions(dataset, filters)
    group_by = str(spec.get("group_by") or "")
    if group_by:
        _group(dataset, group_by)
    metrics = spec.get("metrics") or [{"op": "count", "field": "*"}]
    if not isinstance(metrics, list) or not 1 <= len(metrics) <= 6:
        raise ValueError("Chọn từ 1 đến 6 chỉ số.")
    clean_metrics = []
    for item in metrics:
        if not isinstance(item, dict) or item.get("op") not in NUMERIC_OPS:
            raise ValueError("Phép tổng hợp không được hỗ trợ.")
        op, field = item["op"], item.get("field")
        if field != "*" and field not in dataset.fields:
            raise ValueError("Trường tổng hợp không hợp lệ.")
        if field == "*" and op != "count":
            raise ValueError("Chỉ COUNT được dùng với toàn bộ dòng.")
        if op in {"sum", "avg"} and field not in dataset.numeric:
            raise ValueError("SUM/AVG chỉ dùng cho trường số.")
        if op in {"min", "max"} and field not in dataset.numeric | dataset.dates:
            raise ValueError("MIN/MAX chỉ dùng cho số hoặc ngày.")
        value = {"op": op, "field": field}
        if value not in clean_metrics:
            clean_metrics.append(value)
    sort = spec.get("sort") or {}
    if not isinstance(sort, dict) or sort.get("direction", "desc") not in {"asc", "desc"}:
        raise ValueError("Sắp xếp không hợp lệ.")
    sort_field = sort.get("field") or ("group" if group_by else columns[0])
    metric_keys = {item["op"] + "_" + (item["field"] if item["field"] != "*" else "rows") for item in clean_metrics}
    if sort_field not in (metric_keys | ({"group"} if group_by else set(columns))):
        raise ValueError("Trường sắp xếp không hợp lệ.")
    return {"dataset": dataset.name, "columns": columns, "filters": filters,
            "group_by": group_by, "metrics": clean_metrics,
            "sort": {"field": sort_field, "direction": sort.get("direction", "desc")}}


def _query(spec: dict):
    dataset = DATASETS[spec["dataset"]]
    conditions = _conditions(dataset, spec["filters"])
    group = _group(dataset, spec["group_by"])
    if group is None:
        selected = [dataset.fields[field].label(field) for field in spec["columns"]]
        statement = select(*selected).select_from(dataset.table).where(*conditions)
    else:
        selected = [group]
        for item in spec["metrics"]:
            op, field = item["op"], item["field"]
            column = dataset.fields[field] if field != "*" else None
            aggregate = func.count() if field == "*" else getattr(func, op)(column)
            selected.append(aggregate.label(op + "_" + (field if field != "*" else "rows")))
        statement = select(*selected).select_from(dataset.table).where(*conditions).group_by(group)
    sort_column = next(column for column in selected if column.key == spec["sort"]["field"])
    order = sort_column.asc() if spec["sort"]["direction"] == "asc" else sort_column.desc()
    return statement.order_by(order), selected


def preview_report(raw_spec: dict, *, page: int = 1, page_size: int = 50) -> dict:
    spec = validate_spec(raw_spec)
    page = max(1, min(int(page), 100_000))
    page_size = max(1, min(int(page_size), MAX_PAGE_SIZE))
    query, fields = _query(spec)
    with get_db().engine.connect() as conn:
        total = conn.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
        rows = conn.execute(query.offset((page - 1) * page_size).limit(page_size)).mappings().all()
    return {"items": [{key: _safe_value(row[key]) for key in row.keys()} for row in rows],
            "columns": [column.key for column in fields], "page": page,
            "page_size": page_size, "total": total, "spec": spec}


def _report_record(row) -> dict:
    result = dict(row)
    result["spec"] = {"dataset": result.pop("dataset"), "columns": result.pop("columns_json"),
        "filters": result.pop("filters_json"), "group_by": result.pop("group_by_json"),
        "metrics": result.pop("metrics_json"), "sort": result.pop("sort_json")}
    return result


def list_reports(user_id: str) -> list[dict]:
    with get_db().engine.connect() as conn:
        rows = conn.execute(select(saved_reports).where(saved_reports.c.user_id == user_id)
            .order_by(saved_reports.c.pinned.desc(), saved_reports.c.updated_at.desc())
            .limit(200)).mappings().all()
    return [_report_record(row) for row in rows]


def get_report(user_id: str, report_id: str) -> dict | None:
    with get_db().engine.connect() as conn:
        row = conn.execute(select(saved_reports).where(saved_reports.c.user_id == user_id,
            saved_reports.c.id == report_id)).mappings().first()
    return _report_record(row) if row else None


def save_report(user_id: str, body: dict, report_id: str = "") -> dict:
    spec = validate_spec(body.get("spec"))
    name = str(body.get("name") or "").strip()
    if not 1 <= len(name) <= 120:
        raise ValueError("Tên báo cáo phải có từ 1 đến 120 ký tự.")
    description = str(body.get("description") or "")[:1000]
    if report_id and not get_report(user_id, report_id):
        raise LookupError("Không tìm thấy báo cáo của tài khoản này.")
    stamp = now()
    values = {"name": name, "description": description, "dataset": spec["dataset"],
        "columns_json": spec["columns"], "filters_json": spec["filters"],
        "group_by_json": spec["group_by"], "metrics_json": spec["metrics"],
        "sort_json": spec["sort"], "updated_at": stamp}
    with get_db().engine.begin() as conn:
        if report_id:
            conn.execute(update(saved_reports).where(saved_reports.c.user_id == user_id,
                saved_reports.c.id == report_id).values(**values))
        else:
            report_id = _id()
            conn.execute(insert(saved_reports).values(id=report_id, user_id=user_id,
                pinned=False, created_at=stamp, **values))
    return get_report(user_id, report_id)


def delete_report(user_id: str, report_id: str) -> bool:
    with get_db().engine.begin() as conn:
        owned = conn.scalar(select(saved_reports.c.id).where(saved_reports.c.id == report_id,
            saved_reports.c.user_id == user_id))
        if not owned:
            return False
        conn.execute(update(report_runs).where(report_runs.c.report_id == report_id,
            report_runs.c.user_id == user_id).values(report_id=None))
        conn.execute(delete(saved_reports).where(saved_reports.c.id == report_id,
            saved_reports.c.user_id == user_id))
    return True


def pin_report(user_id: str, report_id: str, pinned: bool) -> dict:
    with get_db().engine.begin() as conn:
        changed = conn.execute(update(saved_reports).where(saved_reports.c.id == report_id,
            saved_reports.c.user_id == user_id).values(pinned=bool(pinned), updated_at=now())).rowcount
    if not changed:
        raise LookupError("Không tìm thấy báo cáo của tài khoản này.")
    return get_report(user_id, report_id)


def run_saved_report(user_id: str, report_id: str, *, page: int = 1, page_size: int = 50) -> dict:
    saved = get_report(user_id, report_id)
    if not saved:
        raise LookupError("Không tìm thấy báo cáo của tài khoản này.")
    started = time.monotonic()
    try:
        result = preview_report(saved["spec"], page=page, page_size=page_size)
    except Exception as exc:
        with get_db().engine.begin() as conn:
            conn.execute(insert(report_runs).values(id=_id(), report_id=report_id,
                user_id=user_id, status="failed", row_count=0,
                duration_ms=int((time.monotonic() - started) * 1000),
                created_at=now(), error_message=str(exc)[:300]))
        raise
    with get_db().engine.begin() as conn:
        conn.execute(insert(report_runs).values(id=_id(), report_id=report_id,
            user_id=user_id, status="completed", row_count=result["total"],
            duration_ms=int((time.monotonic() - started) * 1000),
            created_at=now(), error_message=None))
    return result | {"report": saved}


def list_report_runs(user_id: str, report_id: str = "") -> list[dict]:
    with get_db().engine.connect() as conn:
        conditions = [report_runs.c.user_id == user_id]
        if report_id:
            conditions.append(report_runs.c.report_id == report_id)
        return [dict(row) for row in conn.execute(select(report_runs).where(*conditions)
            .order_by(report_runs.c.created_at.desc()).limit(100)).mappings()]


def _export_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_cell(key: str, value, dataset: ReportDataset):
    if value is not None and key in dataset.dates:
        try:
            raw = str(value)
            return datetime.fromisoformat(raw) if "T" in raw else date.fromisoformat(raw)
        except ValueError:
            pass
    return _export_cell(value)


def export_report(raw_spec: dict, kind: str = "xlsx", *, report_name: str = "",
                  generated_by: str = "") -> tuple[bytes, str]:
    spec = validate_spec(raw_spec)
    query, fields = _query(spec)
    limit = export_sync_max_rows()
    with get_db().engine.connect() as conn:
        rows = conn.execute(query.limit(limit + 1)).mappings().all()
    if len(rows) > limit:
        raise ValueError(f"Báo cáo vượt {limit:,} dòng. Hãy thêm bộ lọc trước khi xuất.")
    headers = [column.key for column in fields]
    filename = "Bao_cao_mua_sam_cong_" + date.today().strftime("%Y%m%d") + "." + kind
    if kind == "csv":
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows([_export_cell(_safe_value(row[key])) for key in headers] for row in rows)
        return ("\ufeff" + stream.getvalue()).encode("utf-8"), filename
    if kind != "xlsx":
        raise ValueError("Định dạng xuất không được hỗ trợ.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Báo cáo"
    sheet.append(headers)
    dataset = DATASETS[spec["dataset"]]
    for row in rows:
        sheet.append([_xlsx_cell(key, row[key], dataset) for key in headers])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1D3448")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(50, max(15, len(key) + 4))
        if key in DATASETS[spec["dataset"]].numeric or key.startswith(("sum_", "avg_")):
            for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.number_format = '#,##0.00'
        elif key in dataset.dates:
            for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.number_format = 'dd/mm/yyyy hh:mm' if any(isinstance(cell.value, datetime) for cell in cells) else 'dd/mm/yyyy'
    info = workbook.create_sheet("Thông tin")
    info.append(["MuaSamCong Pro", "Báo cáo dữ liệu đấu thầu"])
    info.append(["Tên báo cáo", _export_cell(report_name or "Báo cáo tùy chỉnh")])
    info.append(["Dataset", spec["dataset"]])
    info.append(["Bộ lọc", _export_cell(str(spec["filters"]))])
    info.append(["Tạo lúc (UTC)", now()])
    info.append(["Tạo bởi", _export_cell(generated_by or "Tài khoản hiện tại")])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), filename


WIDGET_TYPES = {"kpi", "line", "bar", "table", "changes", "deadlines"}


def _layout(raw) -> list[dict]:
    if not isinstance(raw, list) or len(raw) > 12:
        raise ValueError("Dashboard có tối đa 12 widget.")
    result = []
    used_ids = set()
    for position, item in enumerate(raw):
        if not isinstance(item, dict) or item.get("type") not in WIDGET_TYPES:
            raise ValueError("Loại widget không hợp lệ.")
        spec = validate_spec(item.get("spec"))
        title = str(item.get("title") or "")[:100].strip()
        if not title:
            raise ValueError("Widget cần tên hiển thị.")
        widget_id = str(item.get("id") or "")
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,32}", widget_id) or widget_id in used_ids:
            widget_id = _id()
        used_ids.add(widget_id)
        result.append({"id": widget_id, "type": item["type"],
            "title": title, "spec": spec, "position": position,
            "size": item.get("size") if item.get("size") in {"half", "full"} else "half"})
    return result


def list_dashboards(user_id: str) -> list[dict]:
    with get_db().engine.connect() as conn:
        return [dict(row) for row in conn.execute(select(user_dashboards)
            .where(user_dashboards.c.user_id == user_id)
            .order_by(user_dashboards.c.is_default.desc(), user_dashboards.c.updated_at.desc())
            .limit(100)).mappings()]


def get_dashboard(user_id: str, dashboard_id: str) -> dict | None:
    with get_db().engine.connect() as conn:
        row = conn.execute(select(user_dashboards).where(
            user_dashboards.c.user_id == user_id,
            user_dashboards.c.id == dashboard_id)).mappings().first()
    return dict(row) if row else None


def save_dashboard(user_id: str, body: dict, dashboard_id: str = "") -> dict:
    name = str(body.get("name") or "").strip()
    if not 1 <= len(name) <= 120:
        raise ValueError("Tên dashboard phải có từ 1 đến 120 ký tự.")
    layout = _layout(body.get("layout"))
    if dashboard_id and not get_dashboard(user_id, dashboard_id):
        raise LookupError("Không tìm thấy dashboard của tài khoản này.")
    stamp = now()
    with get_db().engine.begin() as conn:
        if dashboard_id:
            conn.execute(update(user_dashboards).where(user_dashboards.c.id == dashboard_id,
                user_dashboards.c.user_id == user_id).values(
                name=name, layout_json=layout, updated_at=stamp))
        else:
            dashboard_id = _id()
            conn.execute(insert(user_dashboards).values(id=dashboard_id, user_id=user_id,
                name=name, is_default=False, layout_json=layout,
                created_at=stamp, updated_at=stamp))
    return get_dashboard(user_id, dashboard_id)


def set_default_dashboard(user_id: str, dashboard_id: str) -> dict:
    with get_db().engine.begin() as conn:
        owned = conn.scalar(select(user_dashboards.c.id).where(user_dashboards.c.id == dashboard_id,
            user_dashboards.c.user_id == user_id))
        if not owned:
            raise LookupError("Không tìm thấy dashboard của tài khoản này.")
        conn.execute(update(user_dashboards).where(user_dashboards.c.user_id == user_id)
            .values(is_default=False))
        conn.execute(update(user_dashboards).where(user_dashboards.c.id == dashboard_id,
            user_dashboards.c.user_id == user_id).values(is_default=True))
    return get_dashboard(user_id, dashboard_id)


def delete_dashboard(user_id: str, dashboard_id: str) -> bool:
    with get_db().engine.begin() as conn:
        return bool(conn.execute(delete(user_dashboards).where(
            user_dashboards.c.user_id == user_id,
            user_dashboards.c.id == dashboard_id)).rowcount)


def run_dashboard(user_id: str, dashboard_id: str) -> dict:
    dashboard = get_dashboard(user_id, dashboard_id)
    if not dashboard:
        raise LookupError("Không tìm thấy dashboard của tài khoản này.")
    widgets = [{"id": item["id"], "title": item["title"], "type": item["type"],
        "size": item["size"], "data": preview_report(item["spec"], page_size=25)}
        for item in dashboard["layout_json"]]
    return {"dashboard": dashboard, "widgets": widgets, "updated_at": now()}
