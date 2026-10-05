"""Account-scoped Excel/DOCX export of saved AI analyses."""
from __future__ import annotations

import io
import re

from sqlalchemy import select

from server.database import document_analyses, get_db


def _safe(value) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export_analysis(user_id: str, analysis_id: str, format_name: str) -> tuple[bytes, str, str]:
    if not re.fullmatch(r"[0-9a-f]{32}", analysis_id):
        raise LookupError("Không tìm thấy phân tích.")
    with get_db().engine.connect() as conn:
        row = conn.execute(select(document_analyses).where(
            document_analyses.c.id == analysis_id,
            document_analyses.c.user_id == user_id)).mappings().first()
    if not row:
        raise LookupError("Không tìm thấy phân tích.")
    result = row["result_json"]
    name = f"Phan_tich_{row['analysis_type']}_{analysis_id[:8]}"
    if format_name == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        book = Workbook()
        sheet = book.active
        sheet.title = "Phân tích"
        sheet.append(["Loại", "Nội dung", "Giá trị", "Nguồn"])
        if result.get("answer"):
            sheet.append(["Trả lời", "", _safe(result["answer"]),
                          ", ".join(item["label"] for item in result.get("citations", []))])
        for section in result.get("sections", []):
            for item in section.get("items", []):
                sheet.append([_safe(section.get("title")), _safe(item.get("label")),
                    _safe(item.get("value")),
                    ", ".join(source["label"] for source in item.get("citations", []))])
        if result.get("table"):
            products = book.create_sheet("Danh mục")
            products.append(["STT", "Hàng hóa/dịch vụ", "Đơn vị", "Số lượng",
                             "Yêu cầu kỹ thuật", "Xuất xứ", "Giao hàng", "Nguồn"])
            for index, item in enumerate(result["table"], 1):
                products.append([index, _safe(item.get("item")), _safe(item.get("unit")),
                    _safe(item.get("quantity")), _safe(item.get("technical")),
                    _safe(item.get("origin")), _safe(item.get("delivery")),
                    ", ".join(source["label"] for source in item.get("citations", []))])
        if result.get("comparison"):
            compare = book.create_sheet("So sánh")
            compare.append(["Tiêu chí", "Hồ sơ A", "Nguồn A", "Hồ sơ B", "Nguồn B", "Thay đổi"])
            for item in result["comparison"]:
                compare.append([_safe(item.get("criterion")), _safe(item.get("value_a")),
                    ", ".join(source["label"] for source in item.get("citations_a", [])),
                    _safe(item.get("value_b")),
                    ", ".join(source["label"] for source in item.get("citations_b", [])),
                    _safe(item.get("change"))])
        for tab in book:
            tab.freeze_panes = "A2"
            tab.auto_filter.ref = tab.dimensions
            for cell in tab[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="263B61")
            for column in tab.columns:
                letter = column[0].column_letter
                tab.column_dimensions[letter].width = min(62, max(14,
                    max(len(str(cell.value or "")) for cell in column[:100]) + 2))
                for cell in column:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue(), name + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if format_name == "docx":
        from docx import Document
        document = Document()
        document.add_heading("Phân tích hồ sơ đấu thầu", 0)
        document.add_paragraph(f"Loại phân tích: {row['analysis_type']} · Model: {row['model']}")
        if result.get("answer"):
            document.add_heading("Kết quả", 1)
            document.add_paragraph(result["answer"])
            labels = ", ".join(source["label"] for source in result.get("citations", []))
            if labels:
                document.add_paragraph("Nguồn: " + labels, style="Caption")
        for section in result.get("sections", []):
            document.add_heading(section.get("title") or "Nội dung", 1)
            for item in section.get("items", []):
                document.add_paragraph(f"{item.get('label')}: {item.get('value')}")
                labels = ", ".join(source["label"] for source in item.get("citations", []))
                if labels:
                    document.add_paragraph("Nguồn: " + labels, style="Caption")
        if result.get("table"):
            document.add_heading("Danh mục hàng hóa / dịch vụ", 1)
            columns = ("STT", "Hàng hóa/dịch vụ", "Đơn vị", "Số lượng", "Kỹ thuật", "Xuất xứ", "Giao hàng", "Nguồn")
            table = document.add_table(rows=1, cols=len(columns))
            table.style = "Light Shading Accent 1"
            for index, label in enumerate(columns):
                table.rows[0].cells[index].text = label
            for index, item in enumerate(result["table"], 1):
                values = (str(index), item.get("item", ""), item.get("unit", ""),
                          item.get("quantity", ""), item.get("technical", ""),
                          item.get("origin", ""), item.get("delivery", ""),
                          ", ".join(source["label"] for source in item.get("citations", [])))
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    cells[column].text = str(value)
        if result.get("comparison"):
            document.add_heading("So sánh hồ sơ", 1)
            columns = ("Tiêu chí", "Hồ sơ A", "Hồ sơ B", "Thay đổi")
            table = document.add_table(rows=1, cols=len(columns))
            table.style = "Light Shading Accent 1"
            for index, label in enumerate(columns):
                table.rows[0].cells[index].text = label
            for item in result["comparison"]:
                values = (item.get("criterion", ""), item.get("value_a", ""),
                          item.get("value_b", ""), item.get("change", ""))
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    cells[column].text = str(value)
        stream = io.BytesIO()
        document.save(stream)
        return stream.getvalue(), name + ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise ValueError("Định dạng xuất không hợp lệ.")
