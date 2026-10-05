"""Bounded extraction with source locations for private tender documents."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from server.config import max_upload_bytes

MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".txt": "text/plain",
    ".csv": "text/csv",
}
MAX_UNCOMPRESSED = 80 * 1024 * 1024
MAX_CHUNKS = 5000
MAX_CHARS = 1_500_000


@dataclass
class Extraction:
    chunks: list[dict]
    scan_likely: bool = False


def safe_filename(name: str) -> str:
    clean = str(name or "").replace("\\", "/").split("/")[-1]
    clean = re.sub(r"[\x00-\x1f\x7f<>:\"|?*]", "_", clean).strip(" .")
    if not clean or len(clean) > 200 or Path(clean).suffix.lower() not in MIME:
        raise ValueError("Tên hoặc định dạng file không hợp lệ.")
    return clean


def _zip_member_names(data: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED:
                raise ValueError("File nén bên trong quá lớn.")
            return set(archive.namelist())
    except zipfile.BadZipFile as exc:
        raise ValueError("File Office không hợp lệ.") from exc


def validate_file(name: str, content_type: str, data: bytes,
                  *, max_bytes: int | None = None) -> tuple[str, str]:
    filename = safe_filename(name)
    extension = Path(filename).suffix.lower()
    if not data or len(data) > (max_bytes if max_bytes is not None else max_upload_bytes()):
        raise ValueError("File trống hoặc vượt giới hạn dung lượng.")
    declared = str(content_type or "").split(";", 1)[0].strip().lower()
    acceptable = {MIME[extension], "application/octet-stream"}
    if extension in {".csv", ".txt"}:
        acceptable |= {"text/plain", "text/csv", "application/csv"}
    if declared and declared not in acceptable:
        raise ValueError("MIME type không khớp định dạng file.")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("Nội dung không phải PDF hợp lệ.")
    if extension in {".docx", ".xlsx"}:
        names = _zip_member_names(data)
        required = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
        if required not in names:
            raise ValueError("Cấu trúc file Office không hợp lệ.")
    if extension == ".xls" and not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise ValueError("Nội dung không phải XLS hợp lệ.")
    if extension in {".txt", ".csv"} and b"\x00" in data[:4096]:
        raise ValueError("File văn bản chứa dữ liệu nhị phân.")
    return filename, MIME[extension]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1258", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("Không đọc được mã hóa văn bản.")


def _parts(text: str, size: int = 1800):
    clean = re.sub(r"[ \t]+", " ", text).strip()
    for start in range(0, len(clean), size):
        part = clean[start:start + size].strip()
        if part:
            yield part


def extract_text(name: str, data: bytes) -> Extraction:
    extension = Path(name).suffix.lower()
    chunks: list[dict] = []
    total_chars = 0

    def add(text: str, label: str, **location):
        nonlocal total_chars
        for part in _parts(text):
            if len(chunks) >= MAX_CHUNKS or total_chars + len(part) > MAX_CHARS:
                raise ValueError("Tài liệu có quá nhiều nội dung để xử lý trong một lần.")
            chunks.append({"source_id": f"S{len(chunks) + 1:06d}", "label": label,
                           "text": part, **location})
            total_chars += len(part)

    if extension == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if len(reader.pages) > 300:
                raise ValueError("PDF vượt giới hạn 300 trang.")
            blank = 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    blank += 1
                add(text, f"Trang {number}", page_number=number)
            return Extraction(chunks, scan_likely=bool(reader.pages) and blank == len(reader.pages))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Không đọc được cấu trúc PDF.") from exc

    if extension == ".docx":
        from docx import Document
        try:
            document = Document(io.BytesIO(data))
            section = "Mở đầu"
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                    section = text[:200]
                add(text, f"Mục {section}", section=section)
            for table_number, table in enumerate(document.tables, 1):
                for row_number, row in enumerate(table.rows, 1):
                    cells = [cell.text.strip().replace("\n", " · ") for cell in row.cells]
                    add(" | ".join(cells), f"Bảng {table_number}, dòng {row_number}",
                        section=f"Bảng {table_number}", row_start=row_number, row_end=row_number)
            return Extraction(chunks)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Không đọc được cấu trúc DOCX.") from exc

    if extension in {".xlsx", ".xls"}:
        if extension == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                for sheet in book:
                    headers = []
                    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if row_number > 10_000:
                            raise ValueError("Sheet vượt giới hạn 10.000 dòng.")
                        values = [str(value) if value is not None else "" for value in row[:100]]
                        if any(values):
                            if not headers:
                                headers = values
                            normalized = " | ".join(
                                f"{headers[index]}: {value}" if row_number > 1 and index < len(headers)
                                and headers[index] and value else value for index, value in enumerate(values))
                            add(normalized, f'Sheet "{sheet.title}", dòng {row_number}',
                                sheet_name=sheet.title[:120], row_start=row_number, row_end=row_number)
            finally:
                book.close()
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            for sheet in book.sheets():
                headers = []
                if sheet.nrows > 10_000:
                    raise ValueError("Sheet vượt giới hạn 10.000 dòng.")
                for row_number in range(sheet.nrows):
                    values = [str(value) if value is not None else "" for value in sheet.row_values(row_number)[:100]]
                    if any(values):
                        if not headers:
                            headers = values
                        normalized = " | ".join(
                            f"{headers[index]}: {value}" if row_number > 0 and index < len(headers)
                            and headers[index] and value else value for index, value in enumerate(values))
                        add(normalized, f'Sheet "{sheet.name}", dòng {row_number + 1}',
                            sheet_name=sheet.name[:120], row_start=row_number + 1, row_end=row_number + 1)
        return Extraction(chunks)

    decoded = _decode(data)
    if extension == ".csv":
        headers = []
        for row_number, row in enumerate(csv.reader(io.StringIO(decoded)), 1):
            if row_number > 10_000:
                raise ValueError("CSV vượt giới hạn 10.000 dòng.")
            values = row[:100]
            if not headers and any(values):
                headers = values
            normalized = " | ".join(
                f"{headers[index]}: {value}" if row_number > 1 and index < len(headers)
                and headers[index] and value else value for index, value in enumerate(values))
            add(normalized, f"Dòng {row_number}", row_start=row_number, row_end=row_number)
    else:
        for line_number, line in enumerate(decoded.splitlines(), 1):
            add(line, f"Dòng {line_number}", row_start=line_number, row_end=line_number)
    return Extraction(chunks)


class OCRProvider(Protocol):
    def extract(self, name: str, data: bytes) -> Extraction: ...


def ocr_document(name: str, data: bytes, provider: OCRProvider | None = None) -> Extraction:
    if provider is None:
        raise RuntimeError("OCR chưa được cấu hình. PDF scan cần OCR provider riêng.")
    return provider.extract(name, data)
