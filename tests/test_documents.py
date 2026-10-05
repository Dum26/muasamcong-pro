import base64
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from sqlalchemy import func, select

from server.auth import register
from server.database import ai_usage, document_processing, document_text, get_db, reset_db
from server.document_ai import AIUnavailable, analyze, validate_result
from server.document_extract import extract_text, validate_file
from server.document_service import (delete_document, document_content, document_for_user,
                              download_document, process_document, upload_document)


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {
            "DATABASE_URL": "sqlite:///" + (Path(self.folder.name) / "docs.db").as_posix(),
            "ENV": "development", "STORAGE_BACKEND": "local", "MAX_UPLOAD_MB": "10",
            "AI_ENABLED": "true", "AI_PROVIDER": "openai", "AI_MODEL": "test-model",
            "OPENAI_API_KEY": "test-only-no-network"})
        self.env.start()
        reset_db()
        self.user_a = register("docs-a@example.test", "safe-long-password-123")["id"]
        self.user_b = register("docs-b@example.test", "safe-long-password-123")["id"]
        self.created = []

    def tearDown(self):
        for user_id, document_id in self.created:
            delete_document(user_id, document_id)
        reset_db()
        self.env.stop()
        self.folder.cleanup()

    def upload(self, filename, data, mime=""):
        item = upload_document(self.user_a, {"filename": filename, "mime_type": mime,
            "data": base64.b64encode(data).decode()})
        self.created.append((self.user_a, item["id"]))
        return item

    def test_validation_rejects_spoofed_and_oversized_files(self):
        with self.assertRaises(ValueError):
            validate_file("bad.pdf", "application/pdf", b"not pdf")
        with self.assertRaises(ValueError):
            validate_file("bad.txt", "text/plain", b"binary\x00data")
        with self.assertRaises(ValueError):
            validate_file("evil.exe", "application/octet-stream", b"abc")
        with self.assertRaises(ValueError):
            validate_file("big.txt", "text/plain", b"a" * (10 * 1024 * 1024 + 1))
        clean, _ = validate_file("..\\private\\HSMT.txt", "text/plain", "Nội dung".encode())
        self.assertEqual(clean, "HSMT.txt")

    def test_private_upload_extract_ai_cache_and_delete(self):
        item = self.upload("HSMT.txt", "Bảo đảm dự thầu 500 triệu đồng.\nThời gian thực hiện 120 ngày.".encode())
        self.assertIsNone(document_for_user(self.user_b, item["id"]))
        self.assertIsNone(download_document(self.user_b, item["id"]))
        self.assertEqual(process_document(self.user_a, item["id"])["status"], "ready")
        with get_db().engine.connect() as conn:
            attempt = conn.execute(select(document_processing).where(
                document_processing.c.document_id == item["id"])).mappings().one()
        self.assertEqual(attempt["status"], "completed")
        self.assertGreaterEqual(attempt["duration_ms"], 0)
        content = document_content(self.user_a, item["id"])
        self.assertTrue(content["chunks"])
        self.assertIsNone(document_content(self.user_b, item["id"]))
        source = content["chunks"][0]["source_id"]
        response = {"answer": "Bảo đảm dự thầu 500 triệu đồng.", "source_ids": [source],
            "sections": [{"title": "Bảo đảm", "items": [{"label": "Số tiền",
                "value": "500 triệu đồng", "source_ids": [source]}]}], "table": []}
        fake = Mock()
        fake.generate.return_value = response, {"input_tokens": 120, "output_tokens": 50}
        first = analyze(self.user_a, item["id"], "bid_security", provider=fake)
        second = analyze(self.user_a, item["id"], "bid_security", provider=fake)
        self.assertFalse(first["cached"])
        self.assertTrue(second["cached"])
        self.assertEqual(fake.generate.call_count, 1)
        process_document(self.user_a, item["id"])
        third = analyze(self.user_a, item["id"], "bid_security", provider=fake)
        self.assertFalse(third["cached"])
        self.assertEqual(fake.generate.call_count, 2)
        self.assertEqual(first["result"]["sections"][0]["items"][0]["citations"][0]["source_id"], source)
        from server.document_export import export_analysis
        xlsx, _, _ = export_analysis(self.user_a, third["id"], "xlsx")
        docx, _, _ = export_analysis(self.user_a, third["id"], "docx")
        self.assertTrue(xlsx.startswith(b"PK"))
        self.assertTrue(docx.startswith(b"PK"))
        with self.assertRaises(LookupError):
            export_analysis(self.user_b, third["id"], "xlsx")
        with get_db().engine.connect() as conn:
            self.assertEqual(conn.scalar(select(func.count()).select_from(ai_usage)), 2)
        with self.assertRaises(LookupError):
            analyze(self.user_b, item["id"], "summary", provider=fake)
        self.assertFalse(delete_document(self.user_b, item["id"]))
        self.assertTrue(delete_document(self.user_a, item["id"]))
        self.created.clear()
        with get_db().engine.connect() as conn:
            self.assertEqual(conn.scalar(select(func.count()).select_from(document_text)), 0)

    def test_citation_rejects_unknown_source(self):
        with self.assertRaises(AIUnavailable):
            validate_result({"answer": "Có", "source_ids": ["unknown"],
                             "sections": [], "table": []}, [{"source_id": "S000001", "label": "Trang 1"}])

    def test_compare_versions_requires_same_tender_and_scoped_sources(self):
        first = upload_document(self.user_a, {"filename": "v1.txt", "tender_code": "IB2600123456",
            "data": base64.b64encode("Deadline 20/10".encode()).decode()})
        second = upload_document(self.user_a, {"filename": "v2.txt", "tender_code": "IB2600123456",
            "data": base64.b64encode("Deadline 25/10".encode()).decode()})
        self.created += [(self.user_a, first["id"]), (self.user_a, second["id"])]
        process_document(self.user_a, first["id"])
        process_document(self.user_a, second["id"])
        fake = Mock()
        fake.generate.return_value = ({"answer": "Deadline đã đổi.",
            "source_ids": ["A:S000001", "B:S000001"], "sections": [], "table": [],
            "comparison": [{"criterion": "Deadline", "value_a": "20/10", "value_b": "25/10",
                "change": "Lùi 5 ngày", "source_ids_a": ["A:S000001"],
                "source_ids_b": ["B:S000001"]}]}, {})
        result = analyze(self.user_a, first["id"], "compare_versions",
            related_document_id=second["id"], provider=fake)
        self.assertEqual(result["result"]["comparison"][0]["citations_b"][0]["source_id"], "B:S000001")
        with self.assertRaises(LookupError):
            analyze(self.user_b, first["id"], "compare_versions",
                    related_document_id=second["id"], provider=fake)

    def test_docx_and_xlsx_keep_sections_and_rows(self):
        from docx import Document
        from openpyxl import Workbook
        doc = Document()
        doc.add_heading("Năng lực", level=1)
        doc.add_paragraph("Doanh thu tối thiểu 20 tỷ đồng")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Hàng hóa"
        table.cell(0, 1).text = "Số lượng"
        stream = io.BytesIO()
        doc.save(stream)
        item = self.upload("HSMT.docx", stream.getvalue())
        process_document(self.user_a, item["id"])
        chunks = document_content(self.user_a, item["id"])["chunks"]
        self.assertTrue(any(chunk["section"] == "Năng lực" for chunk in chunks))
        self.assertTrue(any("Bảng 1" in chunk["label"] for chunk in chunks))

        book = Workbook()
        sheet = book.active
        sheet.title = "Danh mục"
        sheet.append(["Tên", "Số lượng"])
        sheet.append(["Thuốc", 5])
        stream = io.BytesIO()
        book.save(stream)
        item = self.upload("Phụ lục.xlsx", stream.getvalue())
        process_document(self.user_a, item["id"])
        chunks = document_content(self.user_a, item["id"])["chunks"]
        self.assertTrue(any(chunk["sheet_name"] == "Danh mục" and chunk["row_start"] == 2
                            and "Tên: Thuốc" in chunk["text"] and "Số lượng: 5" in chunk["text"]
                            for chunk in chunks))

    def test_pdf_scan_state_is_explicit(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        stream = io.BytesIO()
        writer.write(stream)
        item = self.upload("scan.pdf", stream.getvalue(), "application/pdf")
        processed = process_document(self.user_a, item["id"])
        self.assertTrue(processed["scan_likely"])
        with self.assertRaisesRegex(ValueError, "OCR"):
            analyze(self.user_a, item["id"], "summary", provider=Mock())


if __name__ == "__main__":
    unittest.main()
