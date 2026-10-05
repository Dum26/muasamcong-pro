import queue
import unittest
import tempfile
from pathlib import Path
from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import Mock, patch

import app


class DateSearchTests(unittest.TestCase):
    def test_plan_search_finds_keyword_in_package_task_summary(self):
        client = app.MuaSamCongClient()
        records = [
            {"id": "named", "planNo": "PL2600000001", "name": "Mua thuốc generic"},
            {"id": "task-1", "planNo": "PL2600299392", "name": "Mua sắm thuốc", "generalTasks": "Thuốc generic"},
            {"id": "task-2", "planNo": "PL2600314291", "name": "Cung cấp thuốc", "generalTasks": "Thuốc generic"},
            {"id": "task-3", "planNo": "PL2600319066", "name": "Mua thuốc", "generalTasks": "Thuốc generic"},
        ]

        def search_response(_url, payload):
            fields = payload[0]["query"][0]["matchFields"]
            matched = [record for record in records if any(
                "generic" in str(record.get(field, "")).lower() for field in fields
            )]
            return {"page": {"content": matched, "totalElements": len(matched)}}

        client._post = Mock(side_effect=search_response)
        found = client.search_plans("generic", max_pages=1,
                                    date_from=date(2026, 9, 11), date_to=date(2026, 9, 24))
        self.assertEqual({record["planNo"] for record in found},
                         {record["planNo"] for record in records})
        self.assertEqual(client.last_search_total, 4)
        query = client._post.call_args.args[1][0]["query"][0]
        self.assertEqual(query["matchFields"], ["planNo", "name", "bidName", "generalTasks"])

    def test_distinct_published_plan_versions_keep_separate_rows_and_codes(self):
        records = [
            {"id": "same-id", "planNo": "PL2600287943", "planNoStand": f"PL2600287943-{version}",
             "planVersion": version, "type": "es-plan-project-p", "name": "Mua thuốc"}
            for version in ("00", "01", "02")
        ]
        candidates = app.MuaSamCongClient._plan_candidates({"content": records})
        rows = [app.plan_row_from_search(record) for record in candidates]
        self.assertEqual(len(candidates), 3)
        self.assertEqual(len(app.group_rows_by_plan(rows, "thuốc")), 3)
        self.assertEqual(len(app.dedupe_rows(rows)), 3)
        self.assertEqual([row.plan_reference for row in rows],
                         ["PL2600287943-00", "PL2600287943-01", "PL2600287943-02"])
        self.assertTrue(all(row.plan_no == "2600287943" for row in rows))

    def test_search_retains_portal_total_when_only_one_page_is_loaded(self):
        client = app.MuaSamCongClient()
        client._smart_search = Mock(return_value={
            "page": {"totalElements": 91, "totalPages": 2},
            "content": [{"id": "one", "planNo": "PL2600287943", "planNoStand": "PL2600287943-02"}],
        })
        self.assertEqual(len(client.search_plans("thuốc", max_pages=1)), 1)
        self.assertEqual(client.last_search_total, 91)

    def test_ninety_one_publications_with_eighty_eight_base_numbers_stay_separate(self):
        records = [
            {"id": f"plan-{i}", "planNo": f"PL2600{i:06d}",
             "planNoStand": f"PL2600{i:06d}-00", "name": "Mua thuốc",
             "publicDate": "2026-09-24T12:00:00"}
            for i in range(88)
        ]
        records += [
            {**records[i], "id": f"revision-{i}",
             "planNoStand": f"PL2600{i:06d}-01"}
            for i in range(3)
        ]
        client = Mock()
        client.search_plans.return_value = records
        worker = SimpleNamespace(
            client=client, keyword_var=Mock(get=Mock(return_value="thuốc")),
            search_mode_var=Mock(get=Mock(return_value="Kế hoạch lựa chọn nhà thầu")),
            pages_var=Mock(get=Mock(return_value=5)),
            detail_var=Mock(get=Mock(return_value=False)),
            fast_mode_var=Mock(get=Mock(return_value=True)),
            notice_status_var=Mock(get=Mock(return_value="Tất cả trạng thái")),
            use_date_filter_var=Mock(get=Mock(return_value=False)),
            remove_dup_var=Mock(get=Mock(return_value=True)),
            msg_queue=queue.Queue(),
        )
        app.App._search_worker(worker)
        events = [worker.msg_queue.get_nowait() for _ in range(worker.msg_queue.qsize())]
        result = next(payload for kind, payload in events if kind == "done")
        self.assertEqual(len(result), 91)
        self.assertIn(("source_count", 91), events)
        self.assertEqual(len({row.plan_reference for row in result}), 91)

    def test_result_id_does_not_imply_a_winner(self):
        for record in ({"inputResultId": "result-1"},
                       {"stepCode": "notify-contractor-step-4-kqlcnt"}):
            with self.subTest(record=record):
                self.assertEqual(app.notice_status_text(record), "Chưa xác định")

    def test_notice_status_takes_priority_over_result_and_document_status(self):
        for code, label in {
            "CNTTT": "Có nhà thầu trúng thầu",
            "KCNTTT": "Không có nhà thầu trúng thầu",
            "DXT": "Đang xét thầu",
            "DHKQLCNT": "Đã hủy KQLCNT",
            "Không công nhận KQLCNT": "Không công nhận KQLCNT",
            "Tuyên bố vô hiệu quyết định về KQLCNT": "Tuyên bố vô hiệu quyết định về KQLCNT",
            "UNKNOWN_STATUS": "UNKNOWN_STATUS",
        }.items():
            with self.subTest(code=code):
                self.assertEqual(app.notice_status_text({
                    "statusForNotify": code, "status": "CANCELED",
                    "inputResultId": "result-1",
                    "stepCode": "notify-contractor-step-4-kqlcnt",
                }), label)

    def test_notice_row_exposes_location_deadline_and_readable_status(self):
        record = {
            "notifyNo": "IB2600548993",
            "notifyNoStand": "IB2600548993-00",
            "planNo": "PL2600314498",
            "bidName": ["Gói thầu thuốc generic"],
            "bidForm": "DTRR",
            "bidCloseDate": "2099-10-29T10:00:00",
            "status": "01",
            "locations": [
                {"provName": "Thành phố Hồ Chí Minh"},
                {"provName": "Thành phố Hồ Chí Minh"},
            ],
        }
        row = app.plan_row_from_search(record)
        self.assertEqual(row.notify_no, "IB2600548993")
        self.assertEqual(row.package_name, "Gói thầu thuốc generic")
        self.assertEqual(row.bid_form, "Đấu thầu rộng rãi")
        self.assertEqual(row.location, "Thành phố Hồ Chí Minh")
        self.assertEqual(row.bid_close_date, "29/10/2099 10:00")
        self.assertEqual(row.tender_status, "Chưa đóng thầu")
        self.assertEqual(app.extract_notify_no({
            "notifyNo": "IB2600000001", "notifyNoStand": "IB2600000001-null",
            "notifyVersion": None,
        }), "IB2600000001")

        now = datetime(2026, 9, 23, 12, 0)
        self.assertEqual(app.notice_status_text({
            "status": "OPEN_BID", "statusForNotify": "DXT",
            "bidCloseDate": "2026-09-20T10:00:00",
        }, now), "Đang xét thầu")
        self.assertEqual(app.notice_status_text({
            "inputResultId": "result-1", "statusForNotify": "CNTTT",
        }, now), "Có nhà thầu trúng thầu")
        self.assertEqual(app.notice_status_text({
            "status": "03", "statusForNotify": "DHTBMT",
        }, now), "Đã hủy TBMT")
        self.assertEqual(app.notice_status_text({
            "status": "CANCEL_BID", "statusForNotify": "DHT",
        }, now), "Đã hủy thầu")
        self.assertEqual(app.notice_status_text({
            "status": "IS_PUBLISH", "statusForNotify": "KCNTTT",
        }, now), "Không có nhà thầu trúng thầu")

    def test_blank_notice_keyword_scans_by_date(self):
        client = app.MuaSamCongClient()
        client._post = Mock(return_value={"page": {"content": [{
            "notifyNo": "IB2600551500", "publicDate": "2026-09-22T23:45:56",
        }]}})
        records = client.search_notices(
            "", max_pages=1, date_from=date(2026, 9, 22), date_to=date(2026, 9, 22)
        )
        self.assertEqual([row["notifyNo"] for row in records], ["IB2600551500"])
        payload = client._post.call_args.args[1][0]
        self.assertEqual(payload["query"][0]["keyWord"], "")
        self.assertIn({
            "fieldName": "publicDate", "searchType": "range",
            "from": "2026-09-22T00:00:00", "to": "2026-09-22T23:59:59.999",
        }, payload["query"][0]["filters"])

    def test_notice_search_excludes_direct_awards_and_result_only_records(self):
        data = {"page": {"content": [
            {
                "notifyNo": "IB2600000001", "stepCode": "notify-contractor-step-1-tbmt",
                "originalPublicDate": "2026-09-17T10:00:00",
            },
            {
                "notifyNo": "DC2600000002", "stepCode": "notify-contractor-step-4-kqlcnt",
                "inputResultId": "direct-result",
            },
            {
                "notifyNo": "IB2600000003", "stepCode": "notify-contractor-step-4-kqlcnt",
                "inputResultId": "result-only", "originalPublicDate": None,
            },
            {
                "notifyNo": "IB2600000004", "stepCode": "notify-contractor-step-4-kqlcnt",
                "inputResultId": "result-with-notice",
                "originalPublicDate": "2026-09-16T08:00:00",
            },
        ]}}
        rows = app.MuaSamCongClient._bid_candidates(data, public_notices_only=True)
        self.assertEqual([row["notifyNo"] for row in rows], [
            "IB2600000001", "IB2600000004",
        ])

    def test_missing_notice_location_is_read_from_plan(self):
        client = app.MuaSamCongClient()
        client.search_plans = Mock(return_value=[{
            "planNo": "PL2600266896",
            "locations": [{
                "provCode": "01", "provName": "Thành phố Hà Nội",
                "districtName": "Xã Đại Thanh",
            }],
        }])
        self.assertEqual(
            client.get_plan_location("2600266896"), "Thành phố Hà Nội"
        )

    def test_bulk_notice_status_filter_skips_lot_lookups(self):
        client = Mock()
        client.get_plan_location.return_value = "Thành phố Hà Nội"
        client.search_notices.return_value = [
            {
                "notifyNo": "IB-OPEN", "planNo": "PL1", "bidName": "Open",
                "bidCloseDate": "2099-09-30T10:00:00", "status": "01",
            },
            {
                "notifyNo": "IB-EVAL", "planNo": "PL2", "bidName": "Evaluate",
                "bidCloseDate": "2026-09-20T10:00:00", "status": "OPEN_BID",
                "statusForNotify": "DXT",
            },
        ]
        worker = Mock(spec=app.App, client=client, msg_queue=queue.Queue())
        for name, value in {
            "keyword_var": "", "search_mode_var": "Thông báo mời thầu",
            "pages_var": 5, "detail_var": False, "use_date_filter_var": True,
            "date_from_var": "01/09/2026", "date_to_var": "23/09/2026",
            "remove_dup_var": True, "notice_status_var": "Đang xét thầu",
        }.items():
            setattr(worker, name, SimpleNamespace(get=lambda value=value: value))
        app.App._search_worker(worker)
        rows = next(payload for kind, payload in worker.msg_queue.queue if kind == "done")
        self.assertEqual([row.notify_no for row in rows], ["IB-EVAL"])
        self.assertEqual(rows[0].location, "Thành phố Hà Nội")
        self.assertEqual(rows[0].lot_file_status, "Không kiểm tra (quét hàng loạt)")
        client.search_plans.assert_not_called()

    def test_fast_plan_scan_skips_per_plan_detail_requests(self):
        client = Mock()
        client.search_plans.return_value = [
            {
                "id": f"plan-{index}", "planNo": f"PL2600{index:06d}",
                "name": f"Kế hoạch generic {index}",
                "publicDate": "2026-09-23T10:00:00",
            }
            for index in range(25)
        ]
        worker = Mock(spec=app.App, client=client, msg_queue=queue.Queue())
        for name, value in {
            "keyword_var": "generic", "search_mode_var": "Kế hoạch lựa chọn nhà thầu",
            "pages_var": 5, "detail_var": False, "fast_mode_var": True,
            "use_date_filter_var": False, "remove_dup_var": True,
        }.items():
            setattr(worker, name, SimpleNamespace(get=lambda value=value: value))

        app.App._search_worker(worker)
        rows = next(payload for kind, payload in worker.msg_queue.queue if kind == "done")

        self.assertEqual(len(rows), 25)
        self.assertTrue(all(row.lot_file_status == "Không kiểm tra (quét nhanh)" for row in rows))
        client.get_plan_detail.assert_not_called()
        client.get_plan_lots.assert_not_called()

    def test_notice_excel_contains_status_location_and_deadline(self):
        from openpyxl import load_workbook

        row = app.ResultRow(
            notify_no="IB2600548993", plan_no="2600314498",
            package_name="Gói thầu thuốc generic", bid_form="Đấu thầu rộng rãi",
            location="Thành phố Hồ Chí Minh", public_date="21/09/2026",
            bid_close_date="29/10/2026 10:00", tender_status="Chưa đóng thầu",
            investor_name="Bệnh viện A", plan_name="Kế hoạch mua thuốc",
            lot_file_status="Không kiểm tra (quét hàng loạt)",
        )
        worker = Mock(spec=app.App, rows=[row], status_var=Mock())
        worker.search_mode_var = SimpleNamespace(get=lambda: "Thông báo mời thầu")
        worker.keyword_var = SimpleNamespace(get=lambda: "")
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "tbmt.xlsx"
            with patch.object(app.filedialog, "asksaveasfilename", return_value=str(target)), \
                    patch.object(app.messagebox, "showinfo"):
                app.App.export_excel(worker)
            wb = load_workbook(target)
            try:
                ws = wb.active
                assert ws is not None
                headers = [cell.value for cell in ws[1]]
                self.assertIn("Mã TBMT", headers)
                self.assertIn("Tỉnh / Thành phố", headers)
                self.assertIn("Thời điểm đóng thầu", headers)
                self.assertIn("Trạng thái", headers)
                values = {ws.cell(1, column).value: ws.cell(2, column).value
                          for column in range(1, ws.max_column + 1)}
                self.assertEqual(values["Mã TBMT"], "IB2600548993")
                self.assertEqual(values["Tỉnh / Thành phố"], "Thành phố Hồ Chí Minh")
                self.assertEqual(values["Trạng thái"], "Chưa đóng thầu")
            finally:
                wb.close()

    def test_plan_search_keeps_only_displayed_package_identity_and_lots(self):
        for get_detail in (False, True):
            with self.subTest(get_detail=get_detail):
                client = Mock()
                client.search_plans.return_value = [{
                    "id": "plan1", "planNo": "PL2600287943", "name": "Plan",
                    "publicDate": "2026-09-09T10:00:00",
                }]
                packages = [
                    {"idDetail": "p1", "bidNo": "BP1", "bidName": "Same name", "bidForm": "DTRR"},
                    {"idDetail": "p2", "bidNo": "BP2", "bidName": "Same name", "bidFormName": "Chỉ định thầu rút gọn"},
                    {"idDetail": "p3", "bidNo": "BP3", "bidName": "No notice", "bidFormName": "Chào hàng cạnh tranh"},
                ]
                client.get_plan_detail.return_value = {"bidpPlanDetailToProjectList": packages}
                client.get_plan_lots.side_effect = [
                    [{"bidNo": "BP1", "lotNo": "PP1"}],
                    [{"bidNo": "BP2", "lotNo": "PP2"}],
                    [],
                ]
                client.search_packages_by_plan.return_value = [
                    {"bidNo": "BP1", "bidName": "Same name", "inputResultId": "result1"},
                ]
                client.get_kqlcnt.return_value = {"bideContractorInputResultDTO": {
                    "planNo": "PL2600287943", "bidName": "Same name",
                    "publicDate": "2026-09-20T10:00:00", "bidFormName": "Other form",
                }}
                worker = Mock(spec=app.App, client=client, msg_queue=queue.Queue())
                for name, value in {
                    "keyword_var": "generic", "search_mode_var": "Kế hoạch lựa chọn nhà thầu",
                    "pages_var": 5, "detail_var": get_detail, "use_date_filter_var": False,
                    "remove_dup_var": True,
                }.items():
                    setattr(worker, name, Mock(get=Mock(return_value=value)))
                app.App._search_worker(worker)
                messages = list(worker.msg_queue.queue)
                self.assertFalse(any(kind == "error" for kind, _ in messages), messages)
                rows = next(payload for kind, payload in messages if kind == "done")
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0].plan_no, "2600287943")
                self.assertEqual(rows[0].package_name, "Same name")
                self.assertEqual(rows[0].bid_form, "Đấu thầu rộng rãi")
                self.assertEqual(rows[0].lot_file_status, "done")
                self.assertEqual([lot["lotNo"] for lot in rows[0].lot_data], ["PP1"])
                self.assertEqual(rows[0].package_no, 'BP1')
                self.assertEqual(rows[0].package_id, 'p1')
                self.assertTrue(rows[0].package_scoped)
                for call, package in zip(client.get_plan_lots.call_args_list, packages):
                    self.assertEqual(call.args[0]["bidpPlanDetailToProjectList"], [package])
                if get_detail:
                    client.get_kqlcnt.assert_called_once_with("result1")
                    self.assertEqual(rows[0].public_date, "20/09/2026")
                else:
                    client.search_packages_by_plan.assert_not_called()

    def test_plan_form_comes_from_matching_package_not_consultancy(self):
        consultancy = app.ResultRow(
            plan_no="2600299008", package_name="Tư vấn lập E-HSMT",
            bid_form="Chỉ định thầu rút gọn", winner="Consultant",
        )
        medicine = app.ResultRow(
            plan_no="2600299008", package_name="Gói số 1: Gói thầu thuốc generic",
            bid_form="Đấu thầu rộng rãi", winner="Medicine supplier",
            lot_file_status="done", lot_data=[{"bidNo": "BP1", "lotNo": "PP1"}],
        )
        rows = app.group_rows_by_plan([consultancy, medicine], "GENERIC")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].bid_form, "Đấu thầu rộng rãi")
        self.assertEqual(rows[0].package_name, medicine.package_name)
        self.assertEqual(rows[0].winner, medicine.winner)
        self.assertEqual(rows[0].lot_data, medicine.lot_data)

    def test_download_selection_keeps_two_packages_from_same_plan(self):
        rows = [app.ResultRow(
            plan_no="2600287943", package_no=bid_no, package_scoped=True,
            lot_file_status="done", lot_data=[{"bidNo": bid_no, "lotNo": "PP1"}],
        ) for bid_no in ("BP1", "BP2")]
        worker = Mock(spec=app.App, rows=rows)
        worker.download_btn = Mock()
        worker.tree = Mock()
        worker.status_var = Mock()
        worker.tree.selection.return_value = ("0", "1")
        with patch.object(app.threading, "Thread") as thread:
            app.App.download_lot_lists(worker)
            self.assertEqual(thread.call_args.kwargs["args"][0], rows)
            self.assertEqual(thread.call_args.kwargs["args"][1], app.LOT_DOWNLOAD_DIR)

    def test_lots_are_scoped_to_plan_packages(self):
        client = app.MuaSamCongClient()
        client.get_package_detail = Mock(return_value={
            "isMultiLot": 1, "bidNo": "BP1", "bidName": "Package 1",
            "bidpBidLotList": [{"lotNo": "PP1", "quantity": 0}],
            "decisionFileId": "not-a-lot-file",
        })
        detail = {"bidpPlanDetailToProjectList": [
            {"idDetail": "p1", "isMultiLot": 1},
            {"idDetail": "p2", "isMultiLot": 0},
        ]}
        lots = client.get_plan_lots(detail)
        self.assertEqual(len(lots), 1)
        self.assertEqual(lots[0]["quantity"], 0)
        client.get_package_detail.assert_called_once_with("p1")
        self.assertEqual(client.get_plan_lots({"bidpPlanDetailToProjectList": []}), [])
        with self.assertRaises(ValueError):
            client.get_plan_lots({})

    def test_lot_export_matches_website_and_does_not_overwrite(self):
        from openpyxl import load_workbook
        row = app.ResultRow(plan_no="2600287943", lot_file_status="done", lot_data=[
            {"bidNo": "BP2600715836-00", "lotNo": "PP1", "lotName": "=NOT_A_FORMULA()",
             "quantity": 0, "pricePlan": 24000, "medicineCode": "G1.01", "medicines": "0",
             "_export_context": {"isMedicine": 1, "planDecisionDate": "2026-09-07"}},
        ])
        with tempfile.TemporaryDirectory() as folder:
            path = app.save_lot_list(row, Path(folder))
            self.assertEqual(path.name, "Danh_sach_phan_lo_BP2600715836-PL2600287943.xlsx")
            second = app.save_lot_list(row, Path(folder))
            self.assertNotEqual(path, second)
            wb = load_workbook(path)
            ws = wb.active
            assert ws is not None
            self.assertEqual(ws.max_row, 2)
            cells = {ws.cell(1, i).value: ws.cell(2, i) for i in range(1, ws.max_column + 1)}
            # Nút Xuất Excel trên web để trống số lượng bằng 0.
            self.assertIsNone(cells["Số lượng"].value)
            self.assertEqual(cells["Tên hoạt chất"].data_type, "s")
            self.assertEqual(cells["Đơn giá (VND)"].value, "24.000")
            self.assertIn("Giá dự toán ước tính từng lô", cells)
            self.assertNotIn("Mã KHLCNT", cells)
            wb.close()

    def test_lot_download_saves_each_package_separately(self):
        from openpyxl import load_workbook

        row = app.ResultRow(plan_no="PL2600287943", lot_file_status="done", lot_data=[
            {"bidNo": "BP2600715836-00", "lotNo": "PP1", "quantity": 1},
            {"bidNo": "BP2600715837-00", "lotNo": "PP2", "quantity": 2},
            {"bidNo": "BP2600715836-00", "lotNo": "PP3", "quantity": 3},
        ])
        for lot in row.lot_data:
            lot["_export_context"] = {"isMedicine": 0, "processApply": "LDT"}
        worker = Mock(spec=app.App, msg_queue=queue.Queue())
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "ketoan"
            app.App._download_lots_worker(worker, [row], destination, 0)
            files = sorted(destination.glob("*.xlsx"))
            self.assertEqual([path.name for path in files], [
                "Danh_sach_phan_lo_BP2600715836-PL2600287943.xlsx",
                "Danh_sach_phan_lo_BP2600715837-PL2600287943.xlsx",
            ])
            for path, expected_lots in zip(files, (["PP1", "PP3"], ["PP2"])):
                wb = load_workbook(path)
                try:
                    values = list(wb.worksheets[0].values)
                    lot_column = values[0].index("Mã phần/lô")
                    self.assertEqual([record[lot_column] for record in values[1:]], expected_lots)
                finally:
                    wb.close()
            kind, summary = worker.msg_queue.get_nowait()
            self.assertEqual(kind, "download_done")
            self.assertIn("Đã lưu 2 file", summary)
            self.assertEqual(len(row.lot_data), 3)

    def test_date_filter_is_sent_on_every_page_in_both_modes(self):
        for method, record_key in (("search_plans", "planNo"),
                                   ("search_notices", "notifyNo")):
            with self.subTest(method=method):
                client = app.MuaSamCongClient()
                first_value = "IB-first" if method == "search_notices" else "first"
                second_value = "IB-second" if method == "search_notices" else "second"
                client._post = Mock(side_effect=[
                    {"page": {"content": [{
                        record_key: first_value,
                        "stepCode": "notify-contractor-step-1-tbmt",
                        "originalPublicDate": "2026-09-09T08:00:00",
                    }]}},
                    {"page": {"content": [{
                        record_key: second_value,
                        "stepCode": "notify-contractor-step-1-tbmt",
                        "originalPublicDate": "2026-09-09T09:00:00",
                    }]}},
                ])
                records = getattr(client, method)(
                    "generic", max_pages=2, page_size=1,
                    date_from=date(2026, 9, 9), date_to=date(2026, 9, 9),
                )
                self.assertEqual(len(records), 2)
                for page, call in enumerate(client._post.call_args_list):
                    payload = call.args[1][0]
                    self.assertEqual(payload["pageNumber"], page)
                    self.assertIn({
                        "fieldName": "publicDate", "searchType": "range",
                        "from": "2026-09-09T00:00:00",
                        "to": "2026-09-09T23:59:59.999",
                    }, payload["query"][0]["filters"])

    def test_no_date_filter_and_one_sided_ranges(self):
        client = app.MuaSamCongClient()
        client._post = Mock(return_value={})
        for start, end in ((None, None), (date(2026, 9, 9), None),
                           (None, date(2026, 9, 9))):
            client._smart_search("generic", 0, 50, "all", ["es-plan-project-p"],
                                 ["name"], date_from=start, date_to=end)
            filters = client._post.call_args.args[1][0]["query"][0]["filters"]
            self.assertEqual(len(filters), 1 if start is end is None else 2)
            if len(filters) == 2:
                self.assertEqual("from" in filters[1], start is not None)
                self.assertEqual("to" in filters[1], end is not None)

    def test_detail_date_does_not_remove_matching_search_result(self):
        client = Mock()
        client.search_notices.return_value = [{
            "planNo": "PL2600297735", "notifyNo": "IB123",
            "publicDate": "2026-09-09T16:55:16", "inputResultId": "result",
        }]
        client.get_kqlcnt.return_value = {"bideContractorInputResultDTO": {
            "planNo": "PL2600297735", "publicDate": "2026-09-20T10:00:00",
        }}
        worker = Mock(spec=app.App, client=client, msg_queue=queue.Queue())
        for name, value in {
            "keyword_var": "generic", "search_mode_var": "Thông báo mời thầu",
            "pages_var": 5, "detail_var": True, "use_date_filter_var": True,
            "date_from_var": "09/09/2026", "date_to_var": "09/09/2026",
            "remove_dup_var": False,
        }.items():
            setattr(worker, name, SimpleNamespace(get=lambda value=value: value))
        app.App._search_worker(worker)
        messages = list(worker.msg_queue.queue)
        self.assertFalse(any(kind == "error" for kind, _ in messages), messages)
        rows = next(payload for kind, payload in messages if kind == "done")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].public_date, "20/09/2026")
        self.assertEqual(client.search_notices.call_args.kwargs["date_from"],
                         date(2026, 9, 9))


if __name__ == "__main__":
    unittest.main()
