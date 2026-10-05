import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app import ResultRow
from web_app import WatchlistStore, QueryHistoryStore, DownloadArchive, build_result_workbook


class WatchlistStoreTest(unittest.TestCase):
    def test_tracking_survives_restart_and_removal(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "watchlist.json"
            store = WatchlistStore(path)
            store.update({"action": "add", "code": "IB2600990001", "name": "Gói thử"})
            self.assertEqual(WatchlistStore(path).read()[0]["name"], "Gói thử")
            WatchlistStore(path).update({"action": "remove", "code": "IB2600990001"})
            self.assertEqual(store.read(), [])

    def test_detects_changes_without_losing_notes_or_creating_false_initial_alert(self):
        with tempfile.TemporaryDirectory() as folder:
            store = WatchlistStore(Path(folder) / "watchlist.json")
            row = {"package_name": "Gói A", "investor_name": "Bệnh viện A",
                   "bid_close_date": "25/09/2026 09:00", "tender_status": "Chưa đóng thầu"}
            first = store.update({"action": "add", "code": "IB2600990001", "row": row})[0]
            self.assertFalse(first["has_update"])
            store.update({"action": "edit", "code": "IB2600990001", "note": "Liên hệ nhóm",
                          "priority": "Cao", "tags": ["y tế"], "pinned": True})
            changed = store.sync([{"code": "IB2600990001", **row,
                                   "bid_close_date": "28/09/2026 09:00", "winner": "Nhà thầu B"}])[0]
            self.assertTrue(changed["has_update"])
            self.assertEqual(changed["note"], "Liên hệ nhóm")
            self.assertEqual(changed["priority"], "Cao")
            self.assertEqual(changed["changes"][0]["field"], "bid_close_date")
            self.assertEqual(len(changed["changes"]), 2)
            self.assertEqual(len(store.sync([{"code": "IB2600990001", **row,
                                              "bid_close_date": "28/09/2026 09:00", "winner": "Nhà thầu B"}])[0]["changes"]), 2)
            store.update({"action": "mark-read", "code": "IB2600990001"})
            self.assertFalse(WatchlistStore(Path(folder) / "watchlist.json").read()[0]["has_update"])

    def test_legacy_watchlist_can_be_edited_and_synced(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "watchlist.json"
            path.write_text('[{"code":"PL2600990001","name":"Kế hoạch cũ","deadline":"25/09/2026"}]', encoding="utf-8")
            store = WatchlistStore(path)
            self.assertEqual(store.read()[0]["workflow_status"], "Mới")
            store.update({"action": "edit", "code": "PL2600990001", "workflow_status": "Đang xem"})
            self.assertEqual(store.read()[0]["workflow_status"], "Đang xem")

    def test_invalid_code_does_not_change_file(self):
        with tempfile.TemporaryDirectory() as folder:
            store = WatchlistStore(Path(folder) / "watchlist.json")
            with self.assertRaises(ValueError):
                store.update({"action": "add", "code": "../../bad"})
            self.assertEqual(store.read(), [])


class QueryHistoryStoreTest(unittest.TestCase):
    def test_recent_queries_survive_restart_and_deduplicate(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "query_history.json"
            store = QueryHistoryStore(path)
            params = {"mode": "Thông báo mời thầu", "keyword": "generic", "pages": 5}
            store.add({"params": params})
            store.add({"params": params})
            self.assertEqual(len(QueryHistoryStore(path).read()), 1)
            self.assertEqual(QueryHistoryStore(path).read()[0]["keyword"], "generic")

    def test_count_pin_rename_delete(self):
        with tempfile.TemporaryDirectory() as folder:
            store = QueryHistoryStore(Path(folder) / "history.json")
            params = {"mode": "Thông báo mời thầu", "keyword": "thiết bị", "pages": 3}
            item = store.add({"params": params})[0]
            store.update({"action": "complete", "params": params, "result_count": 327})
            store.update({"action": "pin", "id": item["id"], "pinned": True})
            store.update({"action": "rename", "id": item["id"], "name": "Thiết bị y tế"})
            self.assertEqual(store.read()[0]["result_count"], 327)
            self.assertEqual(store.read()[0]["name"], "Thiết bị y tế")
            self.assertTrue(store.read()[0]["pinned"])
            store.update({"action": "delete", "id": item["id"]})
            self.assertEqual(store.read(), [])


class DownloadArchiveTest(unittest.TestCase):
    def test_saved_workbook_can_be_downloaded_after_reopening_archive(self):
        with tempfile.TemporaryDirectory() as folder:
            archive = DownloadArchive(Path(folder))
            item = archive.add(b"xlsx bytes", "ket_qua.xlsx", "Excel kết quả")
            reopened = DownloadArchive(Path(folder))
            self.assertEqual(reopened.file(item["id"]), (b"xlsx bytes", "ket_qua.xlsx"))
            self.assertEqual(reopened.read()[0]["size"], 10)
            self.assertIsNone(reopened.file("../../index.json"))


class ResultExportTest(unittest.TestCase):
    def test_lot_column_only_appears_for_plans(self):
        row = ResultRow(notify_no="IB2600123456", plan_no="2600123456",
                        package_name="Gói thử", lot_file_status="done")
        notice = load_workbook(BytesIO(build_result_workbook([row], True)), read_only=True).active
        plan = load_workbook(BytesIO(build_result_workbook([row], False)), read_only=True).active
        self.assertNotIn("Danh sách phân lô", [cell.value for cell in notice[1]])
        self.assertIn("Danh sách phân lô", [cell.value for cell in plan[1]])

if __name__ == "__main__":
    unittest.main()
