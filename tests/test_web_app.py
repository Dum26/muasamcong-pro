import unittest
from io import BytesIO
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from app import ResultRow, official_detail_url, plan_row_from_search
from web_app import (
    JOBS, SearchJob, build_extraction_workbook, build_result_workbook,
    extract_tbmt_codes, row_to_public_dict, selected_rows, start_search,
    EXTRACTIONS, enrich_extraction, lookup_notice_fields,
    load_lots_for_download, prepare_lot_download, build_lot_download,
)


class WebAppTests(unittest.TestCase):
    def test_official_links_keep_the_original_record_identity(self):
        plan = {'id': 'ee8c8976-b9a3-43e8-a276-70eadb5d9380',
                'type': 'es-plan-project-p', 'stepCode': 'plan-step-1',
                'planNo': 'PL2600319066', 'planNoStand': 'PL2600319066-00'}
        url = plan_row_from_search(plan).source_url
        self.assertIn('type=es-plan-project-p', url)
        self.assertIn('id=ee8c8976-b9a3-43e8-a276-70eadb5d9380', url)
        self.assertIn('stepCode=plan-step-1', url)
        notice = {'id': 'aa086665-10d8-4b4e-98df-7e039e6b0a2e',
                  'type': 'es-notify-contractor', 'notifyNo': 'IB2200036956',
                  'notifyId': 'aa086665-10d8-4b4e-98df-7e039e6b0a2e'}
        self.assertIn('notifyNo=IB2200036956', official_detail_url(notice))
        self.assertEqual(official_detail_url({**plan, 'id': 'javascript:alert(1)'}), '')

    def test_job_reports_source_count_separately_from_display_rows(self):
        job = SearchJob({})
        job.update("source_count", 91)
        job.update("done", [ResultRow(plan_no=str(index)) for index in range(88)])
        public = job.public_data()
        self.assertEqual(public["source_count"], 91)
        self.assertEqual(public["count"], 88)

    def test_export_uses_plan_code_without_revision(self):
        row = ResultRow(plan_no="2600287943", plan_reference="PL2600287943-02",
                        plan_name="Mua thuốc")
        sheet = load_workbook(BytesIO(build_result_workbook([row], notice_mode=False))).active
        self.assertEqual(sheet.cell(2, 2).value, "PL2600287943")

    def test_download_fetches_only_the_selected_notice_package(self):
        client = Mock()
        client.search_notices.return_value = [
            {"notifyNo": "IB2600534820", "bidId": "wrong-package"},
            {"notifyNo": "IB2600534827", "planNo": "PL2600275484",
             "bidId": "selected-package"},
        ]
        client.get_package_detail.return_value = {
            "bidNo": "BP1", "bidName": "Selected package", "isMedicine": 1,
            "bidpBidLotList": [{"lotNo": "PP1", "lotName": "Medicine"}],
        }
        with patch('web_app.MuaSamCongClient', return_value=client):
            row = load_lots_for_download(ResultRow(notify_no="IB2600534827"))
        client.get_package_detail.assert_called_once_with("selected-package")
        client.get_plan_lots.assert_not_called()
        self.assertEqual(row.package_no, "BP1")
        self.assertEqual(row.lot_data[0]['bidNo'], "BP1")
        self.assertEqual(row.lot_data[0]['_export_context']['isMedicine'], 1)

    def test_bulk_download_deduplicates_and_reports_unavailable_notices(self):
        rows = [ResultRow(notify_no=code) for code in ('IB1','IB2','IB1')]
        def load(row):
            if row.notify_no == 'IB2':
                raise ValueError('Không có danh sách phân lô')
            return row
        with patch('web_app.load_lots_for_download', side_effect=load) as lookup:
            ready, issues = prepare_lot_download(rows)
        self.assertEqual(lookup.call_count, 2)
        self.assertEqual([r.notify_no for r in ready], ['IB1'])
        self.assertEqual(issues, [{'code':'IB2','reason':'Không có danh sách phân lô'}])
        with self.assertRaises(ValueError):
            prepare_lot_download([])

    def test_each_package_downloads_as_a_separate_excel_file(self):
        rows = [ResultRow(plan_no="2600275484", package_no=bid, lot_file_status="done", lot_data=[
            {"bidNo": bid, "lotNo": "PP1", "lotName": "Medicine", "quantity": 1,
             "_export_context": {"isMedicine": 1, "planDecisionDate": "2026-09-07"}},
        ]) for bid in ('BP1', 'BP2')]
        files, issues = build_lot_download(rows + rows, [{'code': 'IB3', 'reason': 'Missing'}])
        self.assertEqual(len(files), 2)
        self.assertEqual([file['code'] for file in files], ['BP1', 'BP2'])
        self.assertEqual(len({file['name'] for file in files}), 2)
        self.assertEqual(issues, [{'code': 'IB3', 'reason': 'Missing'}])
        for file in files:
            self.assertTrue(file['name'].endswith('.xlsx'))
            self.assertIn(file['code'], file['name'])
            workbook = load_workbook(BytesIO(file['data']))
            self.assertEqual(workbook.sheetnames, ['Danh sách phần lô'])
            self.assertEqual(workbook.active.max_row, 2)
            self.assertEqual(workbook.active['A1'].fill.fgColor.rgb, '00EEEEEE')
            self.assertEqual(workbook.active.column_dimensions['A'].width, 30)

    def test_plan_packages_keep_separate_templates_and_literal_cell_text(self):
        row = ResultRow(plan_no='2600000001', lot_file_status='done', lot_data=[
            {'bidNo': 'BP1', 'lotNo': 'PP1', 'lotName': '=2+2',
             '_export_context': {'isMedicine': 0}},
            {'bidNo': 'BP2', 'lotNo': 'PP2', 'lotName': 'Medicine', 'quantity': 17,
             '_export_context': {'isMedicine': 1, 'planDecisionDate': '2026-09-07'}},
        ])
        with self.assertRaises(ValueError):
            build_lot_download([row], [])  # A plan containing extra packages must not be expanded.
        from dataclasses import replace
        selected = [replace(row, package_no=lot['bidNo'], lot_data=[lot]) for lot in row.lot_data]
        files, issues = build_lot_download(selected, [])
        self.assertEqual(len(files), 2)
        self.assertEqual(issues, [])
        first = load_workbook(BytesIO(files[0]['data'])).active
        second = load_workbook(BytesIO(files[1]['data'])).active
        self.assertNotEqual(first.max_column, second.max_column)
        self.assertEqual(first['C2'].value, '=2+2')
        self.assertEqual(first['C2'].data_type, 's')

    def test_invalid_package_does_not_discard_other_excel_files(self):
        row = ResultRow(plan_no='2600000001', lot_file_status='done', lot_data=[
            {'bidNo': 'BP1', 'lotNo': 'PP1', 'lotName': 'Valid', '_export_context': {}},
            {'bidNo': 'BP2', 'lotNo': 'PP2', 'lotName': 'Missing template'},
        ])
        from dataclasses import replace
        files, issues = build_lot_download([replace(row, package_no=lot['bidNo'], lot_data=[lot]) for lot in row.lot_data], [])
        self.assertEqual([file['code'] for file in files], ['BP1'])
        self.assertEqual(issues[0]['code'], 'BP2')

    def test_individual_links_return_original_bytes_and_expire(self):
        from web_app import LotFileStore
        store = LotFileStore()
        self.addCleanup(store.close)
        with patch('web_app.time.time', return_value=100):
            manifest = store.add([{'data': b'one', 'name': 'BP1.xlsx', 'code': 'BP1'},
                                  {'data': b'two', 'name': 'BP2.xlsx', 'code': 'BP2'}], [])
            tokens = [file['url'].split('=')[1] for file in manifest['files']]
            self.assertEqual(store.read(tokens[0]), (b'one', 'BP1.xlsx'))
            self.assertEqual(store.read(tokens[1]), (b'two', 'BP2.xlsx'))
            self.assertIsNone(store.read('../invalid'))
        with patch('web_app.time.time', return_value=7401):
            self.assertIsNone(store.read(tokens[0]))
            self.assertEqual(store.batches, [])

    def test_pasted_consecutive_codes_use_lookup_fields_not_next_line(self):
        records = extract_tbmt_codes(text="IB2600534827\nIB2600547156\nIB2600534827",
                                     remove_duplicates=False)
        task = {"rows": records, "processed": 0, "state": "running"}
        EXTRACTIONS["lookup-test"] = task
        try:
            with patch('web_app.lookup_notice_fields', return_value={
                "tender_status": "Chưa đóng thầu", "location": "Thành phố Hà Nội",
                "bid_form": "Đấu thầu rộng rãi", "lookup_error": "",
            }) as lookup:
                enrich_extraction("lookup-test")
            self.assertEqual(lookup.call_count, 2)
            self.assertEqual(task["processed"], 3)
            self.assertEqual(task["state"], "done")
            self.assertEqual([r['code'] for r in records],
                             ['IB2600534827','IB2600547156','IB2600534827'])
            sheet = load_workbook(BytesIO(build_extraction_workbook(records))).active
            self.assertEqual(sheet.max_column, 5)
            self.assertEqual(sheet.cell(2, 3).value, "Chưa đóng thầu")
            self.assertEqual(sheet.cell(2, 4).value, "Thành phố Hà Nội")
            self.assertEqual(sheet.cell(2, 5).value, "Đấu thầu rộng rãi")
        finally:
            EXTRACTIONS.pop("lookup-test", None)

    def test_mixed_tbmt_and_khlcnt_codes_use_separate_lookups_and_export(self):
        records = extract_tbmt_codes(text="IB2600534827-00\nPL2600301140-01\npl2600301140")
        self.assertEqual([item["code"] for item in records],
                         ["IB2600534827", "PL2600301140"])
        self.assertEqual([item["code_type"] for item in records], ["TBMT", "KHLCNT"])
        task = {"rows": records, "processed": 0, "state": "running"}
        EXTRACTIONS["mixed-lookup-test"] = task
        try:
            with patch("web_app.lookup_notice_fields", return_value={
                    "tender_status": "Chưa đóng thầu", "title": "Gói A"}) as notice, \
                 patch("web_app.lookup_plan_fields", return_value={
                    "title": "Kế hoạch B", "investor_name": "Bệnh viện B"}) as plan:
                enrich_extraction("mixed-lookup-test")
            notice.assert_called_once_with("IB2600534827")
            plan.assert_called_once_with("PL2600301140")
            self.assertEqual(task["processed"], 2)
            sheet = load_workbook(BytesIO(build_extraction_workbook(records))).active
            self.assertEqual(sheet.cell(1, 3).value, "Mã")
            self.assertEqual(sheet.cell(3, 2).value, "KHLCNT")
            self.assertEqual(sheet.cell(3, 4).value, "Kế hoạch B")
            self.assertEqual(sheet.cell(3, 5).value, "Bệnh viện B")
        finally:
            EXTRACTIONS.pop("mixed-lookup-test", None)

    def test_plan_lookup_uses_exact_pl_code(self):
        from web_app import lookup_plan_fields
        fake = Mock()
        fake.search_plans.return_value = [
            {"planNo": "PL2600301141", "name": "Kế hoạch khác"},
            {"planNo": "PL2600301140", "name": "Kế hoạch đúng",
             "investorName": "Bệnh viện A"},
        ]
        with patch("web_app.MuaSamCongClient", return_value=fake):
            result = lookup_plan_fields("PL2600301140")
        fake.search_plans.assert_called_once_with("PL2600301140", max_pages=1)
        self.assertEqual(result["title"], "Kế hoạch đúng")
        self.assertEqual(result["lookup_error"], "")
        fake.s.close.assert_called_once()

    def test_notice_lookup_requires_exact_code_and_fills_location_from_plan(self):
        client = Mock()
        client.search_notices.return_value = [
            {"notifyNo": "IB2600534820", "statusForNotify": "CNTTT"},
            {"notifyNo": "IB2600534827", "planNo": "PL2600000001",
             "statusForNotify": "DXT", "bidForm": "DTRR"},
        ]
        client.get_plan_location.return_value = "Thành phố Hà Nội"
        with patch('web_app.MuaSamCongClient', return_value=client):
            fields = lookup_notice_fields("IB2600534827")
            self.assertEqual(fields["tender_status"], "Đang xét thầu")
            self.assertEqual(fields["location"], "Thành phố Hà Nội")
            self.assertEqual(fields["bid_form"], "Đấu thầu rộng rãi")
            missing = lookup_notice_fields("IB2600999999")
            self.assertEqual(missing["lookup_error"], "Không tìm thấy TBMT")
            self.assertEqual(missing["tender_status"], "")
            client.search_notices.side_effect = TimeoutError()
            self.assertTrue(lookup_notice_fields("IB2600534827")["lookup_error"])

    def test_public_row_does_not_send_large_lot_payload(self):
        row = ResultRow(notify_no="IB2600000001", lot_file_status="done", lot_data=[{"bidNo": "IB1"}])
        data = row_to_public_dict(row, 3)
        self.assertNotIn("lot_data", data)
        self.assertEqual(data["index"], 3)
        self.assertTrue(data["has_lots"])

    def test_selected_rows_keeps_requested_order_and_removes_duplicates(self):
        job = SearchJob({})
        job.rows = [ResultRow(notify_no="IB1"), ResultRow(notify_no="IB2")]
        rows = selected_rows(job, [1, 1, 0, 99])
        self.assertEqual([row.notify_no for row in rows], ["IB2", "IB1"])

    def test_export_only_86_filtered_rows_from_649_results(self):
        job = SearchJob({})
        job.rows = [ResultRow(notify_no=f"IB2600{index:06d}") for index in range(649)]
        indices = list(range(200, 286))
        rows = selected_rows(job, indices)
        sheet = load_workbook(BytesIO(build_result_workbook(rows, True))).active
        self.assertEqual(sheet.max_row - 1, 86)
        self.assertEqual([sheet.cell(i + 2, 2).value for i in range(86)],
                         [job.rows[i].notify_no for i in indices])
        extracted = extract_tbmt_codes(text="\n".join(row.notify_no for row in rows))
        code_sheet = load_workbook(BytesIO(build_extraction_workbook(extracted))).active
        self.assertEqual(code_sheet.max_row - 1, 86)
        self.assertEqual(code_sheet.cell(87, 1).value, 86)

    def test_empty_export_selection_does_not_export_everything(self):
        job = SearchJob({})
        job.rows = [ResultRow(notify_no="IB1")]
        self.assertEqual(selected_rows(job, []), [])

    def test_web_excel_uses_notice_columns_and_correct_status(self):
        data = build_result_workbook([
            ResultRow(
                notify_no="IB2600000001", plan_no="2600000001",
                package_name="Gói thuốc", tender_status="Chưa đóng thầu",
                location="Thành phố Hà Nội",
            )
        ], notice_mode=True)
        sheet = load_workbook(BytesIO(data)).active
        headers = [cell.value for cell in sheet[1]]
        values = dict(zip(headers, [cell.value for cell in sheet[2]]))
        self.assertEqual(values["Mã TBMT"], "IB2600000001")
        self.assertEqual(values["Trạng thái"], "Chưa đóng thầu")
        self.assertEqual(values["Tỉnh / Thành phố"], "Thành phố Hà Nội")

    def test_completed_query_is_reused_briefly(self):
        params = {"mode": "Thông báo mời thầu", "keyword": "cache-test"}
        existing = JOBS.create(params)
        existing.update("done", [ResultRow(notify_no="IB-CACHED")])
        repeated = start_search(params)
        self.assertEqual(repeated.state, "done")
        self.assertEqual(repeated.rows[0].notify_no, "IB-CACHED")
        self.assertIn("bộ nhớ đệm", repeated.message)

    def test_explicit_refresh_bypasses_recent_search_cache(self):
        params = {"mode": "Thông báo mời thầu", "keyword": "refresh-cache-test"}
        existing = JOBS.create(params)
        existing.update("done", [ResultRow(notify_no="IB-CACHED")])
        with patch("web_app.threading.Thread") as thread:
            fresh = start_search({**params, "force_refresh": True})
        self.assertEqual(fresh.state, "queued")
        self.assertNotIn("force_refresh", fresh.params)
        thread.return_value.start.assert_called_once()

    def test_extracts_normalized_tbmt_codes_and_value_below(self):
        records = extract_tbmt_codes(text=(
            "Mã TBMT : IB2600532133-00\n"
            "Chưa đóng thầu\n"
            "Mã TBMT: IB2600545053-01\n"
            "Đang xét thầu\n"
            "IB2600532133-00"
        ))
        self.assertEqual([item["code"] for item in records], [
            "IB2600532133", "IB2600545053",
        ])
        self.assertEqual(records[0]["cell"], "A1")
        self.assertEqual(records[0]["below_value"], "Chưa đóng thầu")
        self.assertEqual(records[1]["below_value"], "Đang xét thầu")

    def test_extracts_codes_from_xlsx_and_exports_result(self):
        from openpyxl import Workbook, load_workbook

        source = Workbook()
        source.active.title = "test_1"
        source.active["C4"] = "Mã TBMT : IB2600545580-00"
        source.active["C5"] = "Mua thuốc generic"
        buffer = BytesIO()
        source.save(buffer)

        records = extract_tbmt_codes(
            file_data=buffer.getvalue(), filename="danh_sach.xlsx"
        )
        self.assertEqual(records[0]["code"], "IB2600545580")
        self.assertEqual(records[0]["source"], "test_1")
        self.assertEqual(records[0]["cell"], "C4")
        self.assertEqual(records[0]["below_value"], "Mua thuốc generic")

        exported = load_workbook(BytesIO(build_extraction_workbook(records))).active
        values = {exported.cell(1, column).value: exported.cell(2, column).value
                  for column in range(1, exported.max_column + 1)}
        self.assertEqual(values["Mã TBMT"], "IB2600545580")
        self.assertIsNone(values["Trạng thái"])
        self.assertEqual(exported.max_column, 5)


if __name__ == "__main__":
    unittest.main()
