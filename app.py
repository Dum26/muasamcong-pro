# -*- coding: utf-8 -*-
"""
MuaSamCong Tool - Tìm mã KHLCNT theo từ khóa.

Chức năng:
- Tìm "Kế hoạch lựa chọn nhà thầu" theo từ khóa trên muasamcong.mpi.gov.vn.
- Mã KHLCNT hiển thị và sao chép có tiền tố "PL".
- Mã KHLCNT trong nội dung Excel không có tiền tố "PL".
- Có thể lấy thêm KQLCNT: ngày đăng tải, ngày phê duyệt, nhà thầu trúng,
  hình thức LCNT và chủ đầu tư.
- Xuất Excel .xlsx.

Lưu ý:
API của cổng e-GP có thể được thay đổi theo thời gian. Tool cố gắng tự thích nghi
với nhiều dạng JSON và nhiều kiểu matchType. Nếu hệ thống đổi API, xem log trong
file muasamcong_tool.log.
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from io import BytesIO
import logging
import queue
import re
import threading
import time
from dataclasses import dataclass, asdict, field, replace
from datetime import datetime, date, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlencode
from lot_export import EXPORT_CONTEXT_KEYS, lot_export_table

import requests
import ssl
from requests.adapters import HTTPAdapter
from urllib3.util.ssl_ import create_urllib3_context
from server.source.request_manager import get_request_manager
from server.source.errors import BadRequestError
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    from tkcalendar import DateEntry
except ImportError:
    DateEntry = None

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

APP_TITLE = "MuaSamCong Tool - Tra cứu đấu thầu"
BASE = "https://muasamcong.mpi.gov.vn"
SEARCH_API = (
    BASE
    + "/o/egp-portal-contractor-selection-v2/services/smart/search?token=fake"
)
KQLCNT_API = (
    BASE
    + "/o/egp-portal-contractor-selection-v2/services/expose/"
      "contractor-input-result/get?token=fake"
)
PLAN_DETAIL_API = (
    BASE + "/o/egp-portal-contractor-selection-v2/services/expose/lcnt/"
    "bid-po-bidp-plan-project-view/get-by-id?token=fake"
)
PACKAGE_DETAIL_API = (
    BASE + "/o/egp-portal-contractor-selection-v2/services/lcnt/"
    "bid-po-bidp-plan-project-view/get-bidp-plan-detail-by-id?token=fake"
)

if getattr(sys, "frozen", False):
    APP_DIR = Path(sys.executable).resolve().parent
else:
    APP_DIR = Path(__file__).resolve().parent

LOG_FILE = (Path(tempfile.gettempdir()) if os.environ.get("PORT") else APP_DIR) / "muasamcong_tool.log"
LOT_DOWNLOAD_DIR = Path(os.environ.get("MUASAMCONG_DOWNLOAD_DIR") or Path.home() / "Downloads")
logging.basicConfig(
    **({"stream": sys.stderr} if os.environ.get("PORT") else {"filename": LOG_FILE}),
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    encoding="utf-8",
)

HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json; charset=utf-8",
    "Origin": BASE,
    "Referer": BASE + "/web/guest/contractor-selection?render=search",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/153.0.0.0 Safari/537.36"
    ),
}

# Chỉ dùng khi API trả về mã mà không trả về tên hình thức.
# Nếu gặp mã mới, tool vẫn hiển thị nguyên mã thay vì làm mất dữ liệu.
BID_FORM_MAP = {
    "DTRR": "Đấu thầu rộng rãi",
    "DTHC": "Đấu thầu hạn chế",
    "CHCT": "Chào hàng cạnh tranh",
    "CDT": "Chỉ định thầu",
    "CDTRG": "Chỉ định thầu rút gọn",
    "CHCTRG": "Chào hàng cạnh tranh rút gọn",
    "KHAC": "Khác",
    "MSTT": "Mua sắm trực tiếp",
    "TTH": "Tự thực hiện",
    "TVCN": "Tuyển chọn tư vấn cá nhân",
    "DPNG": "Đàm phán giá",
    "CGTT": "Chào giá trực tuyến",
    "CGTTRG": "Chào giá trực tuyến rút gọn",
}


@dataclass
class ResultRow:
    plan_no: str = ""
    plan_reference: str = ""
    notify_no: str = ""
    plan_name: str = ""
    package_name: str = ""
    package_no: str = ""
    package_id: str = ""
    source_url: str = ""
    package_scoped: bool = field(default=False, repr=False)
    public_date: str = ""
    decision_date: str = ""
    winner: str = ""
    bid_form: str = ""
    investor_name: str = ""
    procuring_entity_name: str = ""
    package_price: str = ""
    winning_price: str = ""
    location: str = ""
    bid_close_date: str = ""
    tender_status: str = ""
    lot_file_status: str = "Chưa kiểm tra"
    lot_data: list[dict] = field(default_factory=list, repr=False)


def clean_plan_no(value: Any) -> str:
    """PL2600254687 -> 2600254687."""
    if value is None:
        return ""
    text = str(value).strip().upper()
    m = re.search(r"\bPL\s*(\d+)\b", text)
    if m:
        return m.group(1)
    # Nếu API đã trả số thuần thì giữ nguyên
    m = re.search(r"\b(\d{8,})\b", text)
    return m.group(1) if m else text


def display_plan_no(value: Any) -> str:
    code = clean_plan_no(value)
    return f"PL{code}" if code else ""


def official_detail_url(record: dict) -> str:
    """Link to the published e-GP record, using its original UUID and type."""
    record_id = str(record.get("id") or record.get("notifyId") or "").strip()
    kind = str(record.get("type") or "").strip()
    if not re.fullmatch(r"[0-9a-fA-F]{8}-(?:[0-9a-fA-F]{4}-){3}[0-9a-fA-F]{12}", record_id):
        return ""
    if not (kind.startswith("es-plan-") or kind == "es-notify-contractor"):
        return ""
    step = str(record.get("stepCode") or ("plan-step-1" if kind.startswith("es-plan-") else "notify-contractor-step-1-tbmt"))
    if not re.fullmatch(r"[a-z0-9-]+", step):
        return ""
    namespace = "egpportalcontractorselectionv2_WAR_egpportalcontractorselectionv2"
    params = {
        "p_p_id": namespace, "p_p_lifecycle": "0", "p_p_state": "normal", "p_p_mode": "view",
        f"_{namespace}_render": "detail-v2", "type": kind, "stepCode": step, "id": record_id,
    }
    if kind == "es-notify-contractor":
        for key in ("notifyId", "inputResultId", "bidOpenId", "techReqId", "bidPreNotifyResultId",
                    "bidPreOpenId", "processApply", "bidMode", "notifyNo", "planNo", "bidForm",
                    "isInternet", "caseKHKQ"):
            value = record.get(key)
            params[key] = str(value) if value is not None and value != "" else "undefined"
        params["pno"] = "undefined"
        params["step"] = "tbmt"
    return BASE + "/web/guest/contractor-selection?" + urlencode(params)


def normalize_date(value: Any) -> str:
    """Chuẩn hóa nhiều kiểu ngày của API về dd/mm/YYYY."""
    if value in (None, ""):
        return ""

    if isinstance(value, (int, float)):
        ts = float(value)
        if ts > 10_000_000_000:
            ts /= 1000.0
        try:
            return datetime.fromtimestamp(ts).strftime("%d/%m/%Y")
        except Exception:
            return str(value)

    s = str(value).strip()
    if not s:
        return ""

    # API/UI có thể trả "09/09/2026 - 16:55"
    m = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", s)
    if m:
        dd, mm, yyyy = map(int, m.groups())
        try:
            return date(yyyy, mm, dd).strftime("%d/%m/%Y")
        except ValueError:
            pass

    # hoặc "2026-09-09T16:55:00..."
    m = re.search(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", s)
    if m:
        yyyy, mm, dd = map(int, m.groups())
        try:
            return date(yyyy, mm, dd).strftime("%d/%m/%Y")
        except ValueError:
            pass

    iso = s.replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(iso)
        return dt.strftime("%d/%m/%Y")
    except Exception:
        pass

    return s


def first_value(d: dict, keys: Iterable[str], default: Any = "") -> Any:
    for k in keys:
        if k in d and d[k] not in (None, "", [], {}):
            return d[k]
    return default



def deep_first_value(obj: Any, keys: Iterable[str], default: Any = "") -> Any:
    """
    Tìm field ở mọi cấp JSON.
    Search API e-GP đôi khi đặt dữ liệu trong _source/content/result.
    """
    wanted = {str(k).lower() for k in keys}
    for d in walk_dicts(obj):
        for k, v in d.items():
            if str(k).lower() in wanted and v not in (None, "", [], {}):
                return v
    return default


def text_value(value: Any) -> str:
    """Chuẩn hóa field text; search API đôi khi trả về list một phần tử."""
    if value in (None, ""):
        return ""
    if isinstance(value, (list, tuple, set)):
        values = [text_value(item) for item in value]
        return "; ".join(dict.fromkeys(item for item in values if item))
    if isinstance(value, dict):
        return text_value(first_value(value, ["name", "label", "value"], ""))
    return str(value).strip()


def parse_api_datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        timestamp = float(value)
        if timestamp > 10_000_000_000:
            timestamp /= 1000.0
        try:
            return datetime.fromtimestamp(timestamp)
        except (ValueError, OSError, OverflowError):
            return None
    raw = str(value).strip()
    if not raw:
        return None
    for candidate in (raw.replace("Z", "+00:00"), raw):
        try:
            parsed = datetime.fromisoformat(candidate)
            if parsed.tzinfo is not None:
                return parsed.astimezone(timezone(timedelta(hours=7))).replace(tzinfo=None)
            return parsed
        except ValueError:
            pass
    for fmt in ("%d/%m/%Y - %H:%M", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return None


def normalize_datetime(value: Any) -> str:
    parsed = parse_api_datetime(value)
    return parsed.strftime("%d/%m/%Y %H:%M") if parsed else text_value(value)


def extract_location(record: dict) -> str:
    """Lấy tỉnh/thành phố thực hiện gói thầu, giữ đúng thứ tự từ API."""
    provinces: list[str] = []
    for item in walk_dicts(record):
        locations = item.get("locations")
        if not isinstance(locations, list):
            continue
        for location in locations:
            if not isinstance(location, dict):
                continue
            province = text_value(first_value(
                location,
                ["provName", "provinceName", "cityName", "locationName"],
                "",
            ))
            if province and province not in provinces:
                provinces.append(province)
    if provinces:
        return "; ".join(provinces)
    return text_value(deep_first_value(
        record,
        ["provName", "provinceName", "cityName", "placeOfPerformance"],
        "",
    ))


def extract_notify_no(record: dict) -> str:
    """Trả mã TBMT dạng gọn, không kèm hậu tố phiên bản như -00."""
    base = text_value(first_value(
        record,
        ["notifyNo", "noticeNo"],
        deep_first_value(record, ["notifyNo", "noticeNo"], ""),
    ))
    if base:
        return re.sub(r"-\d{2}$", "", base)
    standard = text_value(record.get("notifyNoStand"))
    standard = re.sub(r"-(?:null|none)$", "", standard, flags=re.IGNORECASE)
    return re.sub(r"-\d{2}$", "", standard)


def notice_status_text(record: dict, now: datetime | None = None) -> str:
    """Đổi mã trạng thái e-GP thành trạng thái dễ đọc cho người dùng."""
    raw_status = text_value(deep_first_value(record, ["status"], "")).upper()
    notify_status = text_value(
        deep_first_value(record, ["statusForNotify"], "")
    ).upper()
    step = text_value(deep_first_value(record, ["stepCode"], "")).lower()
    input_result_id = deep_first_value(record, ["inputResultId"], "")
    close_at = parse_api_datetime(
        deep_first_value(record, ["bidCloseDate", "bidClosingDate", "closeDate"], "")
    )
    now = now or datetime.now()

    status_labels = {
        "CNTTT": "Có nhà thầu trúng thầu",
        "KCNTTT": "Không có nhà thầu trúng thầu",
        "DXT": "Đang xét thầu",
        "DHTBMT": "Đã hủy TBMT",
        "DHT": "Đã hủy thầu",
        "DHKQLCNT": "Đã hủy KQLCNT",
    }
    # Trạng thái thông báo có ưu tiên cao hơn trạng thái tài liệu/ID kết quả.
    if notify_status in status_labels:
        return status_labels[notify_status]
    if notify_status:
        labels = (*status_labels.values(), "Chưa đóng thầu", "Chưa mở thầu",
                  "Không công nhận KQLCNT", "Tuyên bố vô hiệu quyết định về KQLCNT")
        return next((label for label in labels if label.upper() == notify_status),
                    text_value(deep_first_value(record, ["statusForNotify"], "")))
    if raw_status in {"CANCEL_BID", "CANCEL", "CANCELLED", "03"}:
        return "Đã hủy thầu"
    if raw_status == "CANCELED":
        return "Đã hủy KQLCNT"
    if input_result_id or "step-4-kqlcnt" in step:
        return "Chưa xác định"
    if raw_status == "OPEN_BID" or "step-2-kqmt" in step:
        return "Đang xét thầu"
    if close_at is not None:
        return "Chưa đóng thầu" if close_at > now else "Chưa mở thầu"
    return text_value(notify_status or raw_status) or "Chưa xác định"


def extract_public_date(record: dict) -> str:
    value = deep_first_value(
        record,
        [
            "publicDate",
            "publishedDate",
            "publishDate",
            "publicationDate",
            "postingDate",
            "postedDate",
            "createdDate",
            "createDate",
            "publicTime",
            "publishTime",
            "publishedTime",
            "publicDateTime",
            "createdTime",
            "datePublic",
            "datePosted",
        ],
        "",
    )
    return normalize_date(value)


def extract_decision_date(record: dict) -> str:
    value = deep_first_value(
        record,
        [
            "decisionDate",
            "approvalDate",
            "approveDate",
            "approvedDate",
            "decisionTime",
            "approvalTime",
            "approveTime",
        ],
        "",
    )
    return normalize_date(value)


def public_date_sort_key(record: dict) -> float:
    """Giữ cả giờ/phút/giây để xếp bản đăng mới nhất trước như web."""
    value = deep_first_value(record, ["publicDate"], "")
    try:
        if isinstance(value, (int, float)):
            return value / 1000 if value > 10_000_000_000 else float(value)
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (ValueError, TypeError):
        try:
            dt = datetime.strptime(extract_public_date(record), "%d/%m/%Y")
        except (ValueError, TypeError):
            return float("-inf")
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone(timedelta(hours=7)))
    return dt.timestamp()


def walk_dicts(obj: Any):
    """Duyệt mọi dict trong JSON để chịu được thay đổi wrapper response."""
    if isinstance(obj, dict):
        yield obj
        for v in obj.values():
            yield from walk_dicts(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from walk_dicts(v)


def unique_dicts(items: Iterable[dict], keys: list[str]) -> list[dict]:
    out = []
    seen = set()
    for x in items:
        ident = tuple(str(x.get(k, "")) for k in keys)
        if not any(ident):
            ident = (json.dumps(x, ensure_ascii=False, sort_keys=True)[:500],)
        if ident in seen:
            continue
        seen.add(ident)
        out.append(x)
    return out


class LegacySSLAdapter(HTTPAdapter):
    """
    Adapter tương thích máy chủ còn dùng DH/cipher cũ.
    Chỉ áp dụng cho Session của tool này.
    """
    def init_poolmanager(self, connections, maxsize, block=False, **pool_kwargs):
        ctx = create_urllib3_context()
        try:
            ctx.set_ciphers("DEFAULT:@SECLEVEL=1")
        except Exception:
            pass
        pool_kwargs["ssl_context"] = ctx
        return super().init_poolmanager(
            connections, maxsize, block=block, **pool_kwargs
        )



_CLIENT_SESSION_LOCK = threading.Lock()


class MuaSamCongClient:
    def __init__(self, timeout: int = 30, *, force_refresh: bool = False):
        self.timeout = timeout
        self.force_refresh = force_refresh
        self._request_manager = get_request_manager()
        with _CLIENT_SESSION_LOCK:
            if not getattr(self._request_manager, "_legacy_ssl_configured", False):
                self._request_manager.session.headers.update(HEADERS)
                self._request_manager.session.mount("https://", LegacySSLAdapter())
                self._request_manager._legacy_ssl_configured = True
        # Kept for compatibility with existing callers that close ``client.s``.
        # Network traffic itself uses the shared manager/session below.
        self.s = requests.Session()
        self.s.headers.update(HEADERS)
        # Fix lỗi OpenSSL: DH_KEY_TOO_SMALL với một số cấu hình TLS của cổng e-GP.
        self.s.mount("https://", LegacySSLAdapter())

    def _post(self, url: str, payload: Any) -> Any:
        operation = ("search" if url == SEARCH_API else "kqlcnt" if url == KQLCNT_API
                     else "plan" if url == PLAN_DETAIL_API else "package")
        logging.debug("Source request operation=%s", operation)
        ttl_key = ("search_cache_ttl" if operation == "search" else
                   "result_cache_ttl" if operation == "kqlcnt" else
                   "tender_cache_ttl")
        result = self._request_manager.request_json("POST", url,
            operation=operation, payload=payload,
            ttl=self._request_manager.settings[ttl_key], headers=HEADERS,
            timeout=self.timeout, allow_stale=False, force=self.force_refresh)
        if result.stale:
            raise RuntimeError("Nguồn dữ liệu chưa cập nhật; không dùng dữ liệu cũ để ghi đè.")
        return result.data

    def _smart_search(
        self,
        keyword: str,
        page_number: int,
        page_size: int,
        match_type: str,
        field_values: list[str],
        match_fields: list[str],
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> Any:
        filters = [
            {"fieldName": "type", "searchType": "in", "fieldValues": field_values}
        ]
        if date_from is not None or date_to is not None:
            date_filter = {"fieldName": "publicDate", "searchType": "range"}
            if date_from is not None:
                date_filter["from"] = f"{date_from.isoformat()}T00:00:00"
            if date_to is not None:
                date_filter["to"] = f"{date_to.isoformat()}T23:59:59.999"
            filters.append(date_filter)
        payload = [
            {
                "pageSize": page_size,
                "pageNumber": page_number,
                "query": [
                    {
                        "index": "es-contractor-selection",
                        "keyWord": keyword,
                        "matchType": match_type,
                        "matchFields": match_fields,
                        "filters": filters,
                    }
                ],
            }
        ]
        return self._post(SEARCH_API, payload)

    @staticmethod
    def _plan_candidates(data: Any) -> list[dict]:
        candidates = []
        for d in walk_dicts(data):
            t = str(d.get("type", ""))
            if d.get("planNo") or t.startswith("es-plan-"):
                candidates.append(d)
        return unique_dicts(candidates, ["id", "planNo", "planNoStand", "planVersion", "type"])

    @staticmethod
    def _bid_candidates(data: Any, public_notices_only: bool = False) -> list[dict]:
        candidates = []
        for d in walk_dicts(data):
            if d.get("inputResultId") or d.get("notifyNo"):
                if public_notices_only:
                    notify_no = text_value(d.get("notifyNo")).upper()
                    step = text_value(d.get("stepCode")).lower()
                    # Trang TBMT công khai chỉ hiển thị mã IB. Bản step-4 không có
                    # originalPublicDate là bản công bố KQLCNT, không phải TBMT gốc.
                    if not notify_no.startswith("IB"):
                        continue
                    if "step-4-kqlcnt" in step and not d.get("originalPublicDate"):
                        continue
                candidates.append(d)
        keys = ["notifyNo"] if public_notices_only else ["notifyId", "notifyNo", "inputResultId"]
        return unique_dicts(candidates, keys)

    def search_plans(
        self,
        keyword: str,
        max_pages: int = 5,
        page_size: int = 50,
        progress=None,
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> list[dict]:
        """
        Search KHLCNT.

        type 'es-plan-project-p' is the type currently observed on public KHLCNT
        detail/search records. We try several match types because the portal has
        changed keyword matching behavior in past releases.
        """
        keyword = keyword.strip()
        self.last_search_total = None
        if not keyword:
            return []

        # Cùng bốn ô "Tìm trong" của KHLCNT trên e-GP. Thiếu generalTasks sẽ
        # bỏ sót các kế hoạch mà từ khóa chỉ nằm trong tóm tắt gói thầu.
        match_fields = ["planNo", "name", "bidName", "generalTasks"]

        # 'all' is closest to "Khớp tất cả từ"; fallback helps portal revisions.
        match_types = ["all", "any", "contains", "exact"]
        successful_type = None
        first_data = None

        for mt in match_types:
            try:
                data = self._smart_search(
                    keyword=keyword,
                    page_number=0,
                    page_size=page_size,
                    match_type=mt,
                    field_values=["es-plan-project-p"],
                    match_fields=match_fields,
                    date_from=date_from,
                    date_to=date_to,
                )
                plans = self._plan_candidates(data)
                if plans:
                    successful_type = mt
                    first_data = data
                    break
                # Nếu request hợp lệ nhưng chưa có kết quả, tiếp tục thử match khác.
                first_data = data
            except BadRequestError as e:
                logging.warning("matchType %s rejected: %s", mt, e)
                continue

        if successful_type is None:
            # Request có thể hợp lệ nhưng từ khóa thực sự không có dữ liệu.
            return self._plan_candidates(first_data) if first_data is not None else []

        portal_total = (first_data.get("page") or {}).get("totalElements") if isinstance(first_data, dict) else None
        if isinstance(portal_total, (int, float)) and not isinstance(portal_total, bool):
            self.last_search_total = max(0, int(portal_total))
        all_plans = self._plan_candidates(first_data)
        if progress:
            progress(f"Đang quét KHLCNT - trang 1/{max_pages}")

        for page in range(1, max_pages):
            if progress:
                progress(f"Đang quét KHLCNT - trang {page + 1}/{max_pages}")
            data = self._smart_search(
                keyword=keyword,
                page_number=page,
                page_size=page_size,
                match_type=successful_type,
                field_values=["es-plan-project-p"],
                match_fields=match_fields,
                date_from=date_from,
                date_to=date_to,
            )
            batch = self._plan_candidates(data)
            if not batch:
                break
            before = len(all_plans)
            all_plans.extend(batch)
            all_plans = unique_dicts(all_plans, ["id", "planNo", "planNoStand", "planVersion", "type"])
            if len(all_plans) == before:
                break

        return all_plans

    def search_notices(
        self,
        keyword: str,
        max_pages: int = 5,
        page_size: int = 50,
        progress=None,
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> list[dict]:
        """Tìm Thông báo mời thầu."""
        keyword = keyword.strip()

        match_fields = [
            "planNo",
            "notifyNo",
            "bidName",
            "projectName",
            "pName",
            "investorName",
            "procuringEntityName",
        ]
        # Từ khóa trống là chế độ quét hàng loạt theo khoảng ngày.
        match_types = ["all"] if not keyword else ["all", "any", "contains", "exact"]
        successful_type = None
        first_data = None

        for mt in match_types:
            try:
                data = self._smart_search(
                    keyword=keyword,
                    page_number=0,
                    page_size=page_size,
                    match_type=mt,
                    field_values=["es-notify-contractor"],
                    match_fields=match_fields,
                    date_from=date_from,
                    date_to=date_to,
                )
                items = self._bid_candidates(data, public_notices_only=True)
                if items:
                    successful_type = mt
                    first_data = data
                    break
                first_data = data
            except BadRequestError as e:
                logging.warning("notice matchType %s rejected: %s", mt, e)
                continue

        if successful_type is None:
            return self._bid_candidates(first_data, public_notices_only=True) if first_data is not None else []

        all_items = self._bid_candidates(first_data, public_notices_only=True)
        if progress:
            progress(f"Đang quét TBMT - trang 1/{max_pages}")

        for page in range(1, max_pages):
            if progress:
                progress(f"Đang quét TBMT - trang {page + 1}/{max_pages}")
            data = self._smart_search(
                keyword=keyword,
                page_number=page,
                page_size=page_size,
                match_type=successful_type,
                field_values=["es-notify-contractor"],
                match_fields=match_fields,
                date_from=date_from,
                date_to=date_to,
            )
            batch = self._bid_candidates(data, public_notices_only=True)
            if not batch:
                break
            before = len(all_items)
            all_items.extend(batch)
            all_items = unique_dicts(all_items, ["notifyNo"])
            if len(all_items) == before:
                break

        return all_items

    def get_plan_location(self, plan_no: str) -> str:
        """Bổ sung tỉnh/thành từ KHLCNT khi bản ghi TBMT không có locations."""
        code = clean_plan_no(plan_no)
        if not code:
            return ""
        plans = self.search_plans("PL" + code, max_pages=1, page_size=10)
        exact = next(
            (plan for plan in plans if clean_plan_no(plan.get("planNo")) == code),
            None,
        )
        return extract_location(exact) if exact else ""

    def search_packages_by_plan(self, plan_no_with_pl: str) -> list[dict]:
        """Tìm các TBMT/gói thuộc một KHLCNT để lấy inputResultId."""
        plan_no_with_pl = str(plan_no_with_pl).strip().upper()
        if not plan_no_with_pl.startswith("PL"):
            plan_no_with_pl = "PL" + clean_plan_no(plan_no_with_pl)

        candidates = []
        for mt in ("exact", "all", "any"):
            try:
                data = self._smart_search(
                    keyword=plan_no_with_pl,
                    page_number=0,
                    page_size=100,
                    match_type=mt,
                    field_values=["es-notify-contractor"],
                    match_fields=["planNo", "notifyNo", "bidName"],
                )
                batch = self._bid_candidates(data)
                if batch:
                    candidates.extend(batch)
                    break
            except BadRequestError as e:
                logging.warning("package search %s: %s", mt, e)
        return unique_dicts(candidates, ["notifyId", "notifyNo", "inputResultId"])

    def get_kqlcnt(self, input_result_id: str) -> Any:
        return self._post(KQLCNT_API, {"id": input_result_id})

    def get_plan_detail(self, plan_id: str) -> Any:
        return self._post(PLAN_DETAIL_API, {"id": plan_id})

    def get_package_detail(self, package_id: str) -> dict:
        data = self._post(PACKAGE_DETAIL_API, {"id": package_id})
        if not isinstance(data, dict) or (
            "bidpBidLotList" not in data
            and str(data.get("isMultiLot")).lower() not in {"0", "false"}
        ):
            raise ValueError("Không đọc được danh sách phần/lô của gói thầu")
        return data

    def get_plan_lots(self, detail: dict) -> list[dict]:
        packages = detail.get("bidpPlanDetailToProjectList")
        if not isinstance(packages, list):
            raise ValueError("Không đọc được danh sách gói thầu")
        lots = []
        for package in packages:
            if str(package.get("isMultiLot")) == "0":
                continue
            package_id = package.get("idDetail") or package.get("id")
            if not package_id:
                raise ValueError("Gói thầu thiếu mã chi tiết")
            data = self.get_package_detail(str(package_id))
            if str(data.get("isMultiLot")) == "0":
                continue
            export_context = {key: data.get(key) for key in EXPORT_CONTEXT_KEYS}
            for lot in data.get("bidpBidLotList") or []:
                if not isinstance(lot, dict):
                    raise ValueError("Dữ liệu phần/lô không hợp lệ")
                lots.append({**lot, "bidNo": data.get("bidNo") or package.get("bidNo", ""),
                             "bidName": data.get("bidName") or package.get("bidName", ""),
                             "_export_context": export_context})
        return lots


def extract_winners(info: dict) -> str:
    winners = []
    lots = first_value(info, ["lotResultDTO", "lotResults"], [])
    if not isinstance(lots, list):
        lots = []

    # Có trường hợp lotResultDTO nằm ở wrapper bên ngoài info.
    for lot in lots:
        if not isinstance(lot, dict):
            continue
        contractors = first_value(lot, ["contractorList", "contractors"], [])
        if not isinstance(contractors, list):
            continue
        for c in contractors:
            if not isinstance(c, dict):
                continue
            bid_result = c.get("bidResult")
            is_winner = str(bid_result).strip().lower() in {"1", "true", "winner", "win"}
            if is_winner:
                name = first_value(
                    c,
                    ["orgFullname", "contractorName", "orgName", "name"],
                    "",
                )
                if name and str(name).strip() not in winners:
                    winners.append(str(name).strip())

    return "; ".join(winners)


def locate_kqlcnt_info(data: Any) -> dict:
    """
    Tìm dict chứa dữ liệu bideContractorInputResultDTO hoặc dict có các field KQLCNT.
    """
    if isinstance(data, dict):
        dto = data.get("bideContractorInputResultDTO")
        if isinstance(dto, dict):
            # lotResultDTO đôi khi nằm trong DTO
            return dto

    for d in walk_dicts(data):
        dto = d.get("bideContractorInputResultDTO")
        if isinstance(dto, dict):
            return dto

    for d in walk_dicts(data):
        score = sum(
            1
            for k in ("bidName", "publicDate", "decisionDate", "investorName", "bidForm")
            if k in d
        )
        if score >= 3:
            return d
    return {}


def bid_form_text(info: dict) -> str:
    # Ưu tiên tên đã được backend trả về.
    val = first_value(
        info,
        [
            "bidFormName",
            "bidFormNameVi",
            "bidFormNameVN",
            "bidFormText",
            "bidForm",
        ],
        "",
    )
    if not val:
        return ""
    s = text_value(val)
    return BID_FORM_MAP.get(s.upper(), s)


def plan_row_from_search(plan: dict) -> ResultRow:
    plan_raw = deep_first_value(plan, ["planNo", "planNumber"], "")
    plan_reference = text_value(deep_first_value(plan, ["planNoStand"], ""))
    if not plan_reference:
        plan_reference = text_value(plan_raw)
        version = text_value(plan.get("planVersion"))
        if version and plan_reference and not re.search(r"-\d{2}$", plan_reference):
            plan_reference += "-" + version.zfill(2)
    notify_no = extract_notify_no(plan)

    plan_name = deep_first_value(
        plan,
        [
            "planName",
            "name",
            "projectName",
            "pName",
            "subject",
            "planTitle",
        ],
        "",
    )

    package_name = deep_first_value(
        plan,
        ["bidName", "packageName", "bidPackageName"],
        "",
    )

    investor = deep_first_value(
        plan,
        [
            "investorName",
            "procuringEntityName",
            "orgName",
            "investorFullName",
        ],
        "",
    )

    # bidForm_text() chỉ đọc cùng cấp, nên thử tìm sâu nếu cần.
    form = bid_form_text(plan)
    if not form:
        raw_form = deep_first_value(
            plan,
            [
                "bidFormName",
                "bidFormNameVi",
                "bidFormNameVN",
                "bidFormText",
                "bidForm",
                "selectionMethod",
                "selectionMethodName",
            ],
            "",
        )
        if raw_form:
            raw_form = str(raw_form).strip()
            form = BID_FORM_MAP.get(raw_form.upper(), raw_form)

    procuring_entity = deep_first_value(plan, ["procuringEntityName"], "")
    raw_price = plan.get("bidPrice") or plan.get("packagePrice")
    price_unit = text_value(plan.get("bidPriceUnit") or "VND")
    package_price = ""
    if isinstance(raw_price, (int, float)) and not isinstance(raw_price, bool) and raw_price > 0:
        package_price = f"{raw_price:,.0f}".replace(",", ".") + f" {price_unit or 'VND'}"

    return ResultRow(
        plan_no=clean_plan_no(plan_raw),
        plan_reference=plan_reference,
        notify_no=text_value(notify_no),
        plan_name=text_value(plan_name),
        package_name=text_value(package_name),
        package_no=text_value(plan.get("bidNo")),
        package_id=text_value(plan.get("bidId")),
        source_url=official_detail_url(plan),
        public_date=extract_public_date(plan),
        decision_date=extract_decision_date(plan),
        bid_form=form,
        investor_name=text_value(investor),
        procuring_entity_name=text_value(procuring_entity),
        package_price=package_price,
        location=extract_location(plan),
        bid_close_date=normalize_datetime(deep_first_value(
            plan, ["bidCloseDate", "bidClosingDate", "closeDate"], ""
        )),
        tender_status=notice_status_text(plan) if notify_no else "",
    )


def rows_from_kqlcnt(data: Any, fallback_plan: dict) -> list[ResultRow]:
    info = locate_kqlcnt_info(data)
    if not info:
        return []

    raw_plan = first_value(
        info,
        ["planNo", "planNumber"],
        first_value(fallback_plan, ["planNo", "planNumber"], ""),
    )
    plan_name = first_value(
        info,
        ["planName", "projectName", "pName"],
        first_value(
            fallback_plan,
            ["planName", "name", "projectName", "pName"],
            "",
        ),
    )
    package_name = first_value(info, ["bidName", "packageName"], "")
    public_date = normalize_date(first_value(info, ["publicDate", "publishDate"], ""))
    decision_date = normalize_date(first_value(info, ["decisionDate"], ""))
    investor = first_value(
        info,
        ["investorName", "procuringEntityName"],
        first_value(fallback_plan, ["investorName"], ""),
    )
    winner = extract_winners(info)
    form = bid_form_text(info) or bid_form_text(fallback_plan)
    fallback_notice = plan_row_from_search(fallback_plan)

    # Nếu lotResultDTO nằm bên ngoài DTO thì thử ghép vào tạm thời để đọc winner.
    if not winner:
        for d in walk_dicts(data):
            if isinstance(d.get("lotResultDTO"), list):
                temp = dict(info)
                temp["lotResultDTO"] = d["lotResultDTO"]
                winner = extract_winners(temp)
                if winner:
                    break

    return [
        ResultRow(
            plan_no=clean_plan_no(raw_plan),
            notify_no=fallback_notice.notify_no,
            source_url=fallback_notice.source_url,
            plan_name=text_value(plan_name),
            package_name=text_value(package_name),
            public_date=public_date,
            decision_date=decision_date,
            winner=winner,
            bid_form=form,
            investor_name=text_value(investor),
            procuring_entity_name=fallback_notice.procuring_entity_name,
            package_price=fallback_notice.package_price,
            winning_price=text_value(deep_first_value(info,
                ["winningPrice", "winningBidPrice", "awardPrice"], "")),
            location=fallback_notice.location,
            bid_close_date=fallback_notice.bid_close_date,
            tender_status=fallback_notice.tender_status,
        )
    ]


def group_rows_by_plan(rows: list[ResultRow], keyword: str = "") -> list[ResultRow]:
    """Một dòng/KHLCNT, giữ nguyên danh tính và phần/lô của gói đang hiển thị."""
    grouped: dict[str, list[ResultRow]] = {}
    for index, row in enumerate(rows):
        grouped.setdefault(row.plan_reference or row.plan_no or f"missing-plan-{index}", []).append(row)
    result = []
    terms = keyword.casefold().split()
    for members in grouped.values():
        # Ưu tiên gói khớp từ khóa. Nếu tra cứu bằng mã/tên kế hoạch,
        # lấy gói đầu tiên như trang chi tiết KHLCNT mặc định.
        selected = next((row for row in members if terms and all(
            term in row.package_name.casefold() for term in terms
        )), members[0])
        combined = replace(selected, lot_data=list(selected.lot_data))
        if combined.lot_data:
            combined.lot_file_status = "done"
        result.append(combined)
    return result


def dedupe_rows(rows: list[ResultRow]) -> list[ResultRow]:
    seen = set()
    out = []
    for r in rows:
        key = (
            r.notify_no,
            r.plan_no,
            r.plan_reference,
            r.package_no,
            r.package_id,
            r.package_name,
            r.decision_date,
            r.winner,
        )
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


def parse_user_date(value: str):
    """Đọc ngày người dùng theo dd/mm/yyyy hoặc yyyy-mm-dd."""
    value = (value or "").strip()
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError("Ngày phải có dạng dd/mm/yyyy, ví dụ 01/09/2026.")


def row_date_for_filter(row: ResultRow):
    """Ưu tiên ngày đăng tải, nếu thiếu thì dùng ngày phê duyệt."""
    for value in (row.public_date, row.decision_date):
        if not value:
            continue

        normalized = normalize_date(value)
        m = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", normalized)
        if m:
            dd, mm, yyyy = map(int, m.groups())
            try:
                return date(yyyy, mm, dd)
            except ValueError:
                pass

        m = re.search(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", normalized)
        if m:
            yyyy, mm, dd = map(int, m.groups())
            try:
                return date(yyyy, mm, dd)
            except ValueError:
                pass

    return None


def filter_rows_by_date(rows: list[ResultRow], date_from, date_to) -> list[ResultRow]:
    if date_from is None and date_to is None:
        return rows

    out = []
    for row in rows:
        d = row_date_for_filter(row)
        if d is None:
            continue
        if date_from is not None and d < date_from:
            continue
        if date_to is not None and d > date_to:
            continue
        out.append(row)
    return out


def save_lot_list(row: ResultRow, folder: Path) -> Path:
    """Xuất phần/lô của một gói thầu, không ghi đè file đã có."""
    if Workbook is None:
        raise RuntimeError("Chưa cài openpyxl để xuất Excel")
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    if not row.lot_data or row.lot_file_status != "done":
        raise ValueError("Không có danh sách phân lô để tải")
    bid_numbers = {str(lot.get("bidNo") or "").strip() for lot in row.lot_data}
    if len(bid_numbers) != 1 or "" in bid_numbers:
        raise ValueError("Mỗi file phân lô cần có đúng một mã gói thầu")
    bid_no = next(iter(bid_numbers))
    context = row.lot_data[0].get("_export_context")
    if not isinstance(context, dict):
        raise ValueError("Thiếu thông tin mẫu xuất Excel. Hãy tìm kiếm lại trước khi tải.")
    headers, records = lot_export_table(row.lot_data, context)
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Danh sách phần lô"
    edge = Side(style="thin")
    border = Border(top=edge, left=edge, bottom=edge, right=edge)
    ws.append(headers)
    for record in records:
        ws.append(record)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.border = border
        ws.cell(ws.max_row, 1).alignment = Alignment(horizontal="left")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="000000", size=12)
        cell.fill = PatternFill("solid", fgColor="EEEEEE")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 30
    buffer = BytesIO()
    wb.save(buffer)
    code = re.sub(r'[^0-9A-Za-z_-]', '_', clean_plan_no(row.plan_no)) or "khong-ma"
    package_code = re.sub(r'-\d+$', '', bid_no.upper())
    package_code = re.sub(r'[^0-9A-Za-z_-]', '_', package_code)
    stem = f"Danh_sach_phan_lo_{package_code}-PL{code}"
    number = 1
    while True:
        suffix = "" if number == 1 else f" ({number})"
        path = folder / f"{stem}{suffix}.xlsx"
        try:
            with path.open("xb") as output:
                output.write(buffer.getvalue())
            return path
        except FileExistsError:
            number += 1


class App(tk.Tk):
    # Các widget được khởi tạo bởi module ui.
    tree: ttk.Treeview
    keyword_entry: ttk.Entry
    search_btn: ttk.Button
    download_btn: ttk.Button
    progress_bar: ttk.Progressbar
    copy_menu: tk.Menu
    column_menu: tk.Menu
    empty_state: tk.Frame
    empty_title_var: tk.StringVar
    empty_hint_var: tk.StringVar
    selection_var: tk.StringVar

    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        icon_path = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "app_icon.ico"
        if icon_path.exists():
            try:
                self.iconbitmap(str(icon_path))
            except tk.TclError:
                logging.exception("Không thể nạp icon ứng dụng: %s", icon_path)
        width = min(1480, self.winfo_screenwidth() - 60)
        height = min(880, self.winfo_screenheight() - 90)
        self.geometry(f"{width}x{height}+24+24")
        self.minsize(1180, 720)
        self.configure(background="#eef3f8")
        self.search_running = False
        self.copy_cell_value = ""
        self.copy_column_id = "plan_no"

        self.client = MuaSamCongClient()
        self.rows: list[ResultRow] = []
        self.msg_queue: queue.Queue = queue.Queue()

        self.keyword_var = tk.StringVar(value="generic")
        self.search_mode_var = tk.StringVar(value="Kế hoạch lựa chọn nhà thầu")
        self.pages_var = tk.IntVar(value=5)
        self.use_date_filter_var = tk.BooleanVar(value=True)
        self.date_from_var = tk.StringVar(value=date.today().replace(day=1).strftime("%d/%m/%Y"))
        self.date_to_var = tk.StringVar(value=date.today().strftime("%d/%m/%Y"))
        self.detail_var = tk.BooleanVar(value=False)
        self.fast_mode_var = tk.BooleanVar(value=False)
        self.remove_dup_var = tk.BooleanVar(value=True)
        self.notice_status_var = tk.StringVar(value="Tất cả trạng thái")
        self.status_var = tk.StringVar(value="Sẵn sàng")
        self.count_var = tk.StringVar(value="0 kết quả")

        self._build_style()
        self._build_ui()
        self.after(100, self._poll_queue)

    def _build_style(self):
        from ui import build_style
        build_style(self)

    def _build_ui(self):
        from ui import build_ui
        build_ui(self, DateEntry, LOT_DOWNLOAD_DIR)

    def reset(self):
        if self.search_running:
            return
        self.keyword_var.set("")
        self.search_mode_var.set("Kế hoạch lựa chọn nhà thầu")
        self.date_from_var.set(date.today().replace(day=1).strftime("%d/%m/%Y"))
        self.date_to_var.set(date.today().strftime("%d/%m/%Y"))
        self.notice_status_var.set("Tất cả trạng thái")
        self.rows = []
        self._refresh_tree()
        self.status_var.set("Sẵn sàng")

    def _poll_queue(self):
        try:
            while True:
                kind, payload = self.msg_queue.get_nowait()
                if kind == "status":
                    self.status_var.set(payload)
                elif kind == "done":
                    self.search_running = False
                    self.progress_bar.stop()
                    self.rows = payload
                    self.status_var.set(f"Hoàn tất - {len(self.rows)} kết quả")
                    self._refresh_tree()
                    self.search_btn.config(state="normal")
                elif kind == "error":
                    self.search_running = False
                    self.progress_bar.stop()
                    self.search_btn.config(state="normal")
                    self.status_var.set("Có lỗi")
                    self._refresh_tree()
                    messagebox.showerror("Lỗi", payload)
                elif kind == "download_done":
                    self.download_btn.config(state="normal")
                    self.status_var.set(payload)
                    messagebox.showinfo("Tải danh sách phân lô", payload)
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    def start_search(self):
        if self.search_running:
            return
        keyword = self.keyword_var.get().strip()
        mode = self.search_mode_var.get().strip()
        if not keyword and mode != "Thông báo mời thầu":
            messagebox.showwarning("Thiếu từ khóa", "Nhập từ khóa cần tìm.")
            self.keyword_entry.focus_set()
            return

        if not keyword and mode == "Thông báo mời thầu" and not self.use_date_filter_var.get():
            messagebox.showwarning(
                "Chọn khoảng ngày",
                "Khi quét hàng loạt Thông báo mời thầu, hãy bật Lọc theo ngày "
                "để giới hạn dữ liệu cần tải.",
            )
            return

        if self.use_date_filter_var.get():
            try:
                date_from = parse_user_date(self.date_from_var.get())
                date_to = parse_user_date(self.date_to_var.get())
                if date_from and date_to and date_from > date_to:
                    raise ValueError("Từ ngày không được lớn hơn Đến ngày.")
                if not keyword and mode == "Thông báo mời thầu" and not (date_from or date_to):
                    raise ValueError("Quét hàng loạt TBMT cần ít nhất một mốc ngày.")
            except ValueError as e:
                messagebox.showwarning("Ngày không hợp lệ", str(e))
                return

        self.search_btn.config(state="disabled")
        self.search_running = True
        self.progress_bar.start(12)
        self.rows = []
        self._refresh_tree()
        self.status_var.set("Đang kết nối...")

        th = threading.Thread(target=self._search_worker, daemon=True)
        th.start()

    def _search_worker(self):
        started = time.time()
        try:
            keyword = self.keyword_var.get().strip()
            mode = self.search_mode_var.get().strip()
            max_pages = max(1, int(self.pages_var.get()))
            get_detail = bool(self.detail_var.get())
            fast_var = getattr(self, "fast_mode_var", None)
            fast_mode = bool(fast_var.get()) if fast_var is not None else False
            status_var = getattr(self, "notice_status_var", None)
            status_filter = status_var.get().strip() if status_var is not None else "Tất cả trạng thái"

            if self.use_date_filter_var.get():
                date_from = parse_user_date(self.date_from_var.get())
                date_to = parse_user_date(self.date_to_var.get())
            else:
                date_from = None
                date_to = None

            progress = lambda s: self.msg_queue.put(("status", s))

            if mode == "Thông báo mời thầu":
                items = self.client.search_notices(
                    keyword=keyword,
                    max_pages=max_pages,
                    progress=progress,
                    date_from=date_from,
                    date_to=date_to,
                )
            else:
                items = self.client.search_plans(
                    keyword=keyword,
                    max_pages=max_pages,
                    progress=progress,
                    date_from=date_from,
                    date_to=date_to,
                )

            if not items:
                self.msg_queue.put(("done", []))
                return

            # Sắp xếp sau khi gộp các trang, trước khi định dạng mất phần giờ.
            # Chi tiết từng KHLCNT/TBMT vẫn nằm liền nhau theo bản ghi gốc.
            items.sort(key=public_date_sort_key, reverse=True)

            rows: list[ResultRow] = []
            lot_results = {}
            total = len(items)
            portal_total = getattr(self.client, "last_search_total", None) if mode != "Thông báo mời thầu" else None
            source_count = portal_total if isinstance(portal_total, int) and portal_total >= total else total
            # Trạng thái file của TBMT phải kiểm tra theo đúng gói thầu, không
            # dùng chung kết quả phân lô của toàn KHLCNT cho nhiều TBMT.
            bulk_notice = mode == "Thông báo mời thầu"
            bulk_plan_fast = mode != "Thông báo mời thầu" and fast_mode and total > 20
            self.msg_queue.put(
                ("status", f"API tìm thấy {total} bản ghi. Đang xử lý dữ liệu...")
            )
            self.msg_queue.put(("source_count", source_count))

            for i, item in enumerate(items, 1):
                base_row = plan_row_from_search(item)
                if (
                    mode == "Thông báo mời thầu"
                    and status_filter != "Tất cả trạng thái"
                    and base_row.tender_status != status_filter
                ):
                    continue
                if bulk_plan_fast:
                    # Danh sách lớn dùng ngay dữ liệu search API. Người dùng có thể
                    # tắt quét nhanh hoặc tìm riêng một mã để lấy HTLCNT/phân lô.
                    base_row.lot_file_status = "Không kiểm tra (quét nhanh)"
                    rows.append(base_row)
                    continue
                plan_packages = []
                if mode != "Thông báo mời thầu":
                    plan_id = deep_first_value(item, ["id"], "")
                    if plan_id:
                        progress(f"Đang lấy HTLCNT {i}/{total} - {base_row.plan_no}")
                        try:
                            detail = self.client.get_plan_detail(str(plan_id))
                            if not isinstance(detail, dict):
                                raise ValueError("Chi tiết KHLCNT không hợp lệ")
                            plan_packages = detail.get("bidpPlanDetailToProjectList") or []
                            if not isinstance(plan_packages, list) or any(
                                not isinstance(p, dict) for p in plan_packages
                            ):
                                raise ValueError("Danh sách gói thầu không hợp lệ")
                            if plan_packages:
                                notices = []
                                if get_detail:
                                    try:
                                        notices = self.client.search_packages_by_plan("PL" + base_row.plan_no)
                                    except Exception:
                                        logging.exception("Search package notices failed plan=%s", base_row.plan_no)
                                for package in plan_packages:
                                    pkg_row = replace(
                                        base_row,
                                        package_name=str(first_value(package, ["bidName", "packageName", "name"], "")),
                                        package_no=str(package.get("bidNo") or ""),
                                        package_id=str(package.get("idDetail") or package.get("id") or ""),
                                        package_scoped=True,
                                        bid_form=bid_form_text(package),
                                        lot_data=[],
                                    )
                                    try:
                                        pkg_row.lot_data = self.client.get_plan_lots(
                                            {"bidpPlanDetailToProjectList": [package]}
                                        )
                                        pkg_row.lot_file_status = "done" if pkg_row.lot_data else "Ko co file"
                                        if not pkg_row.package_no and pkg_row.lot_data:
                                            pkg_row.package_no = str(pkg_row.lot_data[0].get("bidNo") or "")
                                    except Exception:
                                        logging.exception("Lot list failed package=%s", pkg_row.package_id)
                                        pkg_row.lot_file_status = "Chưa kiểm tra được"
                                    matching = [n for n in notices if pkg_row.package_no
                                                and n.get("bidNo") == pkg_row.package_no]
                                    if not matching and pkg_row.package_name and sum(
                                        str(first_value(p, ["bidName", "packageName", "name"], ""))
                                        == pkg_row.package_name for p in plan_packages
                                    ) == 1:
                                        matching = [n for n in notices if n.get("bidName") == pkg_row.package_name]
                                    # Chỉ bổ sung chi tiết khi xác định được đúng một TBMT.
                                    if len(matching) == 1 and matching[0].get("inputResultId"):
                                        try:
                                            notice = matching[0]
                                            fallback = {**item, **notice, "bidName": pkg_row.package_name}
                                            details = rows_from_kqlcnt(
                                                self.client.get_kqlcnt(str(notice["inputResultId"])), fallback
                                            )
                                            if details:
                                                result = details[0]
                                                pkg_row.public_date = result.public_date or pkg_row.public_date
                                                pkg_row.decision_date = result.decision_date or pkg_row.decision_date
                                                pkg_row.winner = result.winner
                                                pkg_row.investor_name = result.investor_name or pkg_row.investor_name
                                        except Exception:
                                            logging.exception("KQLCNT failed package=%s", pkg_row.package_id)
                                    rows.append(pkg_row)
                                continue
                            try:
                                progress(f"Đang kiểm tra phân lô {i}/{total} - {base_row.plan_no}")
                                base_row.lot_data = self.client.get_plan_lots(detail)
                                base_row.lot_file_status = "done" if base_row.lot_data else "Ko co file"
                            except Exception:
                                logging.exception("Lot list failed planId=%s", plan_id)
                                base_row.lot_file_status = "Chưa kiểm tra được"
                        except Exception:
                            logging.exception("HTLCNT failed planId=%s", plan_id)
                            base_row.bid_form = "Chưa lấy được HTLCNT"
                            base_row.lot_file_status = "Chưa kiểm tra được"
                elif bulk_notice:
                    # Web kiểm tra từng TBMT theo mã gói sau khi trả danh sách.
                    base_row.lot_file_status = "Không kiểm tra (quét hàng loạt)"
                elif base_row.plan_no:
                    if base_row.plan_no in lot_results:
                        base_row.lot_file_status, base_row.lot_data = lot_results[base_row.plan_no]
                    else:
                        try:
                            plans = self.client.search_plans("PL" + base_row.plan_no, max_pages=1)
                            plan = next(p for p in plans if clean_plan_no(p.get("planNo")) == base_row.plan_no)
                            base_row.lot_data = self.client.get_plan_lots(self.client.get_plan_detail(plan["id"]))
                            base_row.lot_file_status = "done" if base_row.lot_data else "Ko co file"
                        except Exception:
                            logging.exception("Lot lookup failed plan=%s", base_row.plan_no)
                            base_row.lot_file_status = "Chưa kiểm tra được"
                lot_results[base_row.plan_no] = (base_row.lot_file_status, base_row.lot_data)

                # HTLCNT luôn được lấy, tùy chọn chỉ điều khiển chi tiết KQLCNT.
                if not get_detail:
                    rows.append(base_row)
                    continue

                display_code = base_row.plan_no or clean_plan_no(item.get("planNo", ""))
                self.msg_queue.put(
                    ("status", f"Đang lấy chi tiết {i}/{total} - {display_code}")
                )

                if mode == "Thông báo mời thầu":
                    # TBMT thường có inputResultId ngay trên record nếu đã có KQLCNT.
                    input_id = item.get("inputResultId")
                    if input_id:
                        try:
                            data = self.client.get_kqlcnt(str(input_id))
                            detail_rows = rows_from_kqlcnt(data, item)
                            if detail_rows:
                                rows.extend(detail_rows)
                                continue
                        except Exception:
                            logging.exception(
                                "KQLCNT failed inputResultId=%s", input_id
                            )

                    # Chưa có KQLCNT: vẫn giữ Hình thức LCNT từ TBMT.
                    rows.append(base_row)
                    continue

                # Mode KHLCNT: tìm các TBMT thuộc mã kế hoạch rồi lấy KQLCNT.
                raw_plan_no = first_value(item, ["planNo", "planNumber"], "")
                if not raw_plan_no:
                    rows.append(base_row)
                    continue

                try:
                    packages = self.client.search_packages_by_plan(str(raw_plan_no))
                except Exception:
                    logging.exception("Search packages failed for %s", raw_plan_no)
                    rows.append(base_row)
                    continue

                detail_added = False
                for pkg in packages:
                    # Gộp dữ liệu KHLCNT + TBMT để fallback được bidForm,
                    # investorName, bidName nếu response KQLCNT thiếu field.
                    fallback = dict(item)
                    fallback.update({k: v for k, v in pkg.items() if v not in (None, "", [], {})})
                    if not bid_form_text(fallback):
                        matching_forms = {
                            bid_form_text(p) for p in plan_packages
                            if isinstance(p, dict) and p.get("bidName") == pkg.get("bidName")
                            and p.get("bidName") and bid_form_text(p)
                        }
                        if len(matching_forms) == 1:
                            fallback["bidFormName"] = matching_forms.pop()

                    input_id = pkg.get("inputResultId")
                    if input_id:
                        try:
                            data = self.client.get_kqlcnt(str(input_id))
                            detail_rows = rows_from_kqlcnt(data, fallback)
                            if detail_rows:
                                rows.extend(detail_rows)
                                detail_added = True
                                continue
                        except Exception:
                            logging.exception(
                                "KQLCNT failed inputResultId=%s", input_id
                            )

                    # Dù chưa có KQLCNT, vẫn thêm TBMT để hiện Hình thức LCNT.
                    pkg_row = plan_row_from_search(fallback)
                    if pkg_row.package_name or pkg_row.bid_form:
                        rows.append(pkg_row)
                        detail_added = True

                if not detail_added:
                    rows.append(base_row)

            if mode == "Thông báo mời thầu":
                missing_locations = sorted({
                    row.plan_no for row in rows if row.plan_no and not row.location
                })
                if missing_locations:
                    progress(
                        f"Đang bổ sung tỉnh/thành cho {len(missing_locations)} kế hoạch..."
                    )
                    timeout = getattr(self.client, "timeout", 30)
                    if not isinstance(timeout, (int, float)):
                        timeout = 30

                    def lookup_location(plan_no: str) -> tuple[str, str]:
                        if isinstance(self.client, MuaSamCongClient):
                            client = MuaSamCongClient(timeout=int(timeout))
                            return plan_no, client.get_plan_location(plan_no)
                        value = self.client.get_plan_location(plan_no)
                        return plan_no, value if isinstance(value, str) else ""

                    locations_by_plan: dict[str, str] = {}
                    workers = min(5, len(missing_locations))
                    with ThreadPoolExecutor(max_workers=workers) as executor:
                        futures = {
                            executor.submit(lookup_location, plan_no): plan_no
                            for plan_no in missing_locations
                        }
                        for completed, future in enumerate(as_completed(futures), 1):
                            plan_no = futures[future]
                            try:
                                code, location = future.result()
                                if location:
                                    locations_by_plan[code] = location
                            except Exception:
                                logging.exception(
                                    "Plan location lookup failed plan=%s", plan_no
                                )
                            progress(
                                f"Đang bổ sung tỉnh/thành {completed}/{len(futures)}"
                            )
                    for row in rows:
                        if not row.location:
                            row.location = locations_by_plan.get(row.plan_no, "")

            # API lọc ngày đăng tải KHLCNT/TBMT trước khi phân trang.
            # Không lọc lại theo ngày KQLCNT: đó là ngày của tài liệu khác.
            for row in rows:
                if row.package_scoped:
                    continue
                if row.lot_file_status == "Không kiểm tra (quét nhanh)":
                    continue
                row.lot_file_status, row.lot_data = lot_results.get(
                    row.plan_no, ("Chưa kiểm tra được", [])
                )

            if mode != "Thông báo mời thầu":
                rows = group_rows_by_plan(rows, keyword)

            if self.remove_dup_var.get():
                rows = dedupe_rows(rows)

            elapsed = time.time() - started
            self.msg_queue.put(("status", f"Hoàn tất trong {elapsed:.1f} giây"))
            self.msg_queue.put(("done", rows))

        except requests.RequestException as e:
            logging.exception("Network error")
            self.msg_queue.put(
                (
                    "error",
                    "Không kết nối được API Mua Sắm Công.\n\n"
                    f"{e}\n\n"
                    "Kiểm tra Internet, thử mở trang muasamcong.mpi.gov.vn trên trình duyệt "
                    "và chạy lại tool.",
                )
            )
        except Exception as e:
            logging.exception("Unhandled error")
            self.msg_queue.put(
                (
                    "error",
                    f"{type(e).__name__}: {e}\n\n"
                    f"Chi tiết đã ghi tại:\n{LOG_FILE}",
                )
            )

    def _refresh_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        for i, r in enumerate(self.rows, 1):
            self.tree.insert(
                "",
                "end",
                iid=str(i - 1),
                tags=("alternate",) if i % 2 == 0 else (),
                values=(
                    i,
                    r.notify_no,
                    display_plan_no(r.plan_no),
                    r.plan_name,
                    r.package_name,
                    r.public_date,
                    r.bid_close_date,
                    r.tender_status,
                    r.location,
                    r.decision_date,
                    r.winner,
                    r.bid_form,
                    r.investor_name,
                    r.lot_file_status,
                ),
            )
        self.count_var.set(f"{len(self.rows)} kết quả")
        if "empty_state" in self.__dict__:
            if self.rows:
                self.empty_state.place_forget()
            else:
                self.empty_state.place(relx=.5, rely=.54, anchor="center")
                if self.search_running:
                    title, hint = "Đang tìm kiếm…", "Thông tin kế hoạch và phân lô đang được cập nhật."
                elif self.status_var.get() == "Có lỗi":
                    title, hint = "Chưa tải được dữ liệu", "Kiểm tra kết nối rồi thử tìm kiếm lại."
                elif self.keyword_var.get().strip():
                    title, hint = "Chưa tìm thấy kết quả", "Thử từ khóa khác hoặc mở rộng khoảng ngày tìm kiếm."
                elif self.search_mode_var.get() == "Thông báo mời thầu":
                    if self.status_var.get().startswith("Hoàn tất"):
                        title, hint = "Chưa tìm thấy TBMT", "Thử mở rộng khoảng ngày hoặc chọn trạng thái khác."
                    else:
                        title, hint = "Quét TBMT theo khoảng ngày", "Có thể để trống từ khóa để lấy hàng loạt thông báo mời thầu."
                else:
                    title, hint = "Bắt đầu với một từ khóa", "Nhập tên gói thầu hoặc mã KHLCNT rồi nhấn Tìm kiếm."
                self.empty_title_var.set(title)
                self.empty_hint_var.set(hint)
            self.selection_var.set("Chưa chọn dòng")

    def _put_clipboard(self, text, status):
        if not text:
            self.status_var.set("Ô được chọn chưa có dữ liệu")
            return "break"
        self.clipboard_clear()
        self.clipboard_append(text)
        self.status_var.set(status)
        return "break"

    def download_lot_lists(self):
        if str(self.download_btn.cget("state")) == "disabled":
            return
        selected = set(self.tree.selection())
        if not selected:
            self.status_var.set("Chọn một hoặc nhiều dòng rồi bấm Tải phân lô")
            return
        chosen = {}
        skipped = 0
        for i, row in enumerate(self.rows):
            if str(i) not in selected:
                continue
            if row.lot_file_status == "done" and row.lot_data:
                chosen.setdefault((row.plan_no, row.package_no, row.package_id, row.package_name), row)
            else:
                skipped += 1
        if not chosen:
            messagebox.showinfo("Danh sách phân lô", "Các dòng đã chọn không có danh sách phân lô để tải.\nNếu trạng thái là Chưa kiểm tra được, hãy tìm kiếm lại.")
            return
        folder = LOT_DOWNLOAD_DIR
        self.download_btn.config(state="disabled")
        self.status_var.set("Đang lưu danh sách phân lô...")
        threading.Thread(target=self._download_lots_worker,
                         args=(list(chosen.values()), Path(folder), skipped), daemon=True).start()

    def _download_lots_worker(self, rows, folder, skipped):
        saved, errors = [], []
        try:
            folder.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            self.msg_queue.put(("download_done", f"Không tạo được thư mục lưu {folder}: {exc}"))
            return
        for row in rows:
            packages: dict[str, list[dict]] = {}
            for lot in row.lot_data:
                bid_no = str(lot.get("bidNo") or "").strip()
                packages.setdefault(bid_no, []).append(lot)
            for bid_no, lots in packages.items():
                try:
                    saved.append(save_lot_list(replace(row, lot_data=lots), folder))
                except Exception as exc:
                    logging.exception("Save lots failed plan=%s bid=%s", row.plan_no, bid_no)
                    errors.append(f"{row.plan_no} / {bid_no}: {exc}")
        summary = f"Đã lưu {len(saved)} file tại:\n{folder}"
        if skipped:
            summary += f"\nBỏ qua {skipped} dòng không có dữ liệu tải."
        if errors:
            summary += "\nKhông lưu được:\n" + "\n".join(errors)
        self.msg_queue.put(("download_done", summary))

    def _select_all_rows(self, event=None):
        self.tree.selection_set(self.tree.get_children())
        return "break"

    def _cell_at(self, event):
        row = self.tree.identify_row(event.y)
        col = self.tree.identify_column(event.x)
        if self.tree.identify_region(event.x, event.y) != "cell" or not row or not col:
            return None
        return row, self.tree.set(row, col)

    def _copy_clicked_cell(self, event):
        cell = self._cell_at(event)
        if cell:
            return self._put_clipboard(cell[1], "Đã sao chép nội dung ô")

    def _show_copy_menu(self, event):
        region = self.tree.identify_region(event.x, event.y)
        col = self.tree.identify_column(event.x)
        if region not in {"cell", "heading"} or not col:
            return
        self.copy_column_id = self.tree.column(col, "id")
        title = self.tree.heading(self.copy_column_id, "text")
        menu = self.column_menu if region == "heading" else self.copy_menu
        menu.entryconfigure(0, label=f"Sao chép cả cột: {title}")
        if region == "cell":
            cell = self._cell_at(event)
            if cell is None:
                return
            row, self.copy_cell_value = cell
            if row not in self.tree.selection():
                self.tree.selection_set(row)
        self.tree.focus_set()
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def copy_column_values(self, column=None, selected_only=False):
        column = column or self.copy_column_id
        self.copy_column_id = column
        selected = set(self.tree.selection())
        items = [item for item in self.tree.get_children()
                 if not selected_only or item in selected]
        if not items:
            self.status_var.set("Chọn dòng cần sao chép trước" if selected_only else "Chưa có dữ liệu để sao chép")
            return "break"
        # Giữ cả ô trống và giá trị trùng để dán các cột không bị lệch dòng.
        values = [str(self.tree.set(item, column)).replace("\t", " ")
                  .replace("\r", " ").replace("\n", " ") for item in items]
        self.clipboard_clear()
        self.clipboard_append("\n".join(values))
        title = self.tree.heading(column, "text")
        self.status_var.set(f"Đã sao chép cột {title} • {len(items)} dòng • Dán bằng Ctrl+V")
        return "break"

    def copy_selected_rows(self):
        selected = set(self.tree.selection())
        if not selected:
            self.status_var.set("Chọn dòng cần sao chép trước")
            return "break"
        columns = self.tree["displaycolumns"]
        lines = []
        for item in self.tree.get_children():
            if item in selected:
                lines.append("\t".join(str(self.tree.set(item, c)).replace("\t", " ").replace("\r", " ").replace("\n", " ") for c in columns))
        return self._put_clipboard("\n".join(lines), f"Đã sao chép {len(lines)} dòng • Có thể dán vào Excel")

    def copy_codes(self, selected_only=False):
        codes = []
        seen = set()
        selected = set(self.tree.selection())
        source = [r for i, r in enumerate(self.rows) if str(i) in selected] if selected_only else self.rows
        for r in source:
            code = display_plan_no(r.plan_no)
            if code and code not in seen:
                seen.add(code)
                codes.append(code)

        if not codes:
            self.status_var.set("Chọn dòng cần sao chép mã" if selected_only else "Chưa có mã KHLCNT để sao chép")
            return

        self._put_clipboard("\n".join(codes), f"Đã sao chép {len(codes)} mã KHLCNT")

    def export_excel(self):
        if not self.rows:
            messagebox.showinfo("Xuất Excel", "Chưa có dữ liệu để xuất.")
            return

        if Workbook is None:
            messagebox.showerror(
                "Thiếu thư viện",
                "Chưa cài openpyxl.\nChạy: pip install -r requirements.txt",
            )
            return

        notice_mode = self.search_mode_var.get() == "Thông báo mời thầu"
        keyword_part = self.keyword_var.get().strip() or "Tat_ca"
        prefix = "TBMT" if notice_mode else "KHLCNT"
        path = filedialog.asksaveasfilename(
            title="Lưu file Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile=f"{prefix}_{keyword_part}_{datetime.now():%Y%m%d_%H%M}.xlsx",
        )
        if not path:
            return

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.worksheets[0]
        ws.title = "Kết quả"

        if notice_mode:
            headers = [
                "STT", "Mã TBMT", "Mã KHLCNT", "Tên gói thầu",
                "Hình thức LCNT", "Tỉnh / Thành phố", "Ngày đăng tải",
                "Thời điểm đóng thầu", "Trạng thái", "Chủ đầu tư",
                "Tên KHLCNT / Dự án", "Nhà thầu trúng", "Ngày phê duyệt",
                "Danh sách phân lô",
            ]
            widths = [7, 19, 16, 48, 24, 24, 16, 21, 22, 42, 42, 38, 16, 28]
        else:
            headers = [
                "STT", "Mã KHLCNT", "Tên KHLCNT / Dự án", "Tên gói thầu",
                "Ngày đăng tải", "Ngày phê duyệt", "Nhà thầu trúng",
                "Hình thức LCNT", "Chủ đầu tư", "Danh sách phân lô",
            ]
            widths = [7, 16, 38, 42, 16, 16, 40, 24, 40, 24]
        ws.append(headers)

        for i, r in enumerate(self.rows, 1):
            if notice_mode:
                values = [
                    i, r.notify_no, r.plan_no, r.package_name, r.bid_form,
                    r.location, r.public_date, r.bid_close_date, r.tender_status,
                    r.investor_name, r.plan_name, r.winner, r.decision_date,
                    r.lot_file_status,
                ]
            else:
                values = [
                    i, r.plan_no, r.plan_name, r.package_name, r.public_date,
                    r.decision_date, r.winner, r.bid_form, r.investor_name,
                    r.lot_file_status,
                ]
            ws.append(values)

        header_fill = PatternFill("solid", fgColor="0B2239")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        code_columns = (2, 3) if notice_mode else (2,)
        alternate_fill = PatternFill("solid", fgColor="F3F7FA")
        for row in range(2, ws.max_row + 1):
            for column in code_columns:
                ws.cell(row=row, column=column).number_format = "@"
            if row % 2 == 0:
                for cell in ws[row]:
                    cell.fill = alternate_fill
            for cell in ws[row]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
        self.status_var.set(f"Đã xuất: {path}")
        messagebox.showinfo("Xuất Excel", "Xuất Excel thành công.")


if __name__ == "__main__":
    App().mainloop()
